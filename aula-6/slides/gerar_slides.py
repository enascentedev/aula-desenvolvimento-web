# -*- coding: utf-8 -*-
"""
Gera os slides da Aula 06 - A API de login (Python + FastAPI, uv e venv).

Estilo: dark tech / editor de codigo — o mesmo sistema visual das Aulas 04 e 05.
Conteudo: o back-end prometido no fim da aula 5. Primeiro os conceitos
(cliente-servidor, API, HTTP, JSON, FastAPI, venv, uv), depois o passo a passo
para criar, rodar e testar no Swagger uma API com uma rota so: POST /login.

Uso:
    python gerar_slides.py

Saida:
    Aula06-API-Login-FastAPI.pptx  (editavel no PowerPoint)

Para mudar o visual, edite a secao PALETA / TIPOGRAFIA logo abaixo.
Para mudar o conteudo, edite a secao CONTEUDO DOS SLIDES no final do arquivo.
"""

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.parts.image import Image as PptxImage
from pptx.util import Inches, Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


# ---------------------------------------------------------------- PALETA ----

BG = RGBColor(0x12, 0x12, 0x19)          # fundo
BG_HI = RGBColor(0x1D, 0x1A, 0x2E)       # topo do gradiente de fundo
SURFACE = RGBColor(0x1E, 0x1E, 0x2A)     # cartoes e painel de codigo
SURFACE_2 = RGBColor(0x24, 0x24, 0x32)   # barra de janela do painel
BORDER = RGBColor(0x2E, 0x2E, 0x3E)      # bordas
TEXT = RGBColor(0xE8, 0xE8, 0xF0)        # texto principal
TEXT_DIM = RGBColor(0x9A, 0x9A, 0xB0)    # texto secundario
TEXT_FAINT = RGBColor(0x6A, 0x6A, 0x80)  # texto apagado / codigo fora de foco
ACCENT = RGBColor(0x7C, 0x6C, 0xF5)      # indigo (mesma familia do #4f46e5 do CSS da aula)
ACCENT_DEEP = RGBColor(0x3B, 0x34, 0x70) # indigo escuro (decoracao)
ACCENT_2 = RGBColor(0x4E, 0xC9, 0xF0)    # ciano (tags JSX)
OK = RGBColor(0x4E, 0xC9, 0xA0)          # verde
ERR = RGBColor(0xF0, 0x71, 0x78)         # vermelho
HILITE = RGBColor(0x26, 0x22, 0x4A)      # faixa de destaque atras da linha em foco

# cores de sintaxe (familia VS Code Dark+)
C_KEYWORD = RGBColor(0xC5, 0x86, 0xC0)
C_STRING = RGBColor(0xCE, 0x91, 0x78)
C_FUNC = RGBColor(0xDC, 0xDC, 0xAA)
C_TAG = RGBColor(0x4E, 0xC9, 0xF0)
C_VAR = RGBColor(0x9C, 0xDC, 0xFE)
C_NUM = RGBColor(0xB5, 0xCE, 0xA8)
C_PUNCT = RGBColor(0xC8, 0xC8, 0xD4)
C_COMMENT = RGBColor(0x6A, 0x6A, 0x80)
C_SELECTOR = RGBColor(0xD7, 0xBA, 0x7D)


# ------------------------------------------------------------ TIPOGRAFIA ----

# Display: Bahnschrift condensada (geometrica, tecnica) — nada de Segoe UI
# generico nos titulos. Corpo: Segoe UI, escolhido pela legibilidade no
# projetor. Codigo: Cascadia Mono, a fonte do VS Code / Windows Terminal —
# vem do mundo do proprio assunto da aula.
DISPLAY = "Bahnschrift SemiBold Condensed"
SANS = "Segoe UI"
SANS_LIGHT = "Segoe UI Light"
MONO = "Cascadia Mono"

CODE_SIZE = Pt(13)
CODE_LH = Pt(19)          # altura exata de cada linha de codigo (usada no destaque)
CODE_LH_IN = 19 / 72.0    # a mesma coisa, em polegadas
DUPLO_SIZE = Pt(10.5)     # o arquivo inteiro em duas colunas
DUPLO_LH = Pt(14.5)
TERM_SIZE = Pt(12)        # painel de terminal, que ocupa a largura toda
TERM_LH = Pt(17)
PAINEL_LARGO = 7.45       # largura do painel de codigo nos slides de codigo


# ---------------------------------------------------------------- CANVAS ----

W = 13.333   # polegadas
H = 7.5
M = 0.62     # margem lateral
TOTAL_SLIDES = 45
AULA = "Aula 06"


# -------------------------------------------------------- INSTITUCIONAL ----

ESCOLA = "Escola Técnica Mesquita"
PROFESSOR = "Emanuel Nascente"

# O logo e arte preta/vermelha/amarela sobre branco: nao fecha contraste
# sobre o dark do deck, e recortar o fundo apagaria o contorno preto do
# livro. Em vez disso a marca vira um SELO — placa clara com o mesmo raio
# das pilulas do deck — para pertencer ao sistema visual em vez de parecer
# colada nele. De quebra e o unico acento quente (vermelho/amarelo) num
# deck frio de indigo e ciano.
ASSETS = Path(__file__).with_name("assets")
LOGO = ASSETS / "logo-escola-tecnica-mesquita.png"
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # layout em branco
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    # camada de profundidade: gradiente diagonal sutil sobre o fundo chapado
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                Inches(W), Inches(H))
    sp.line.fill.background()
    sp.shadow.inherit = False
    gradient(sp, BG_HI, BG, angle=315)
    return slide


def gradient(shape, cor_a, cor_b, angle=90):
    """Preenche a shape com um gradiente linear de cor_a para cor_b."""
    shape.fill.gradient()
    stops = shape.fill.gradient_stops
    stops[0].color.rgb = cor_a
    stops[0].position = 0.0
    stops[1].color.rgb = cor_b
    stops[1].position = 1.0
    for extra in list(stops)[2:]:
        extra.color.rgb = cor_b
        extra.position = 1.0
    shape.fill.gradient_angle = angle
    return shape


# --------------------------------------------------------------- HELPERS ----

def rect(slide, x, y, w, h, fill=None, line=None, line_w=Pt(1), radius=None,
         shape=MSO_SHAPE.RECTANGLE):
    """Retangulo (opcionalmente arredondado) em polegadas."""
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        sp.adjustments[0] = radius
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w
    sp.shadow.inherit = False
    sp.text_frame.text = ""
    return sp


def oval(slide, x, y, w, h, fill=None, line=None, line_w=Pt(1), rotation=0):
    sp = rect(slide, x, y, w, h, fill=fill, line=line, line_w=line_w,
              shape=MSO_SHAPE.OVAL)
    sp.rotation = rotation
    return sp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def style_run(run, text, size, color, font=SANS, bold=False, italic=False):
    run.text = text
    run.font.size = size
    run.font.color.rgb = color
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic
    return run


def para(tf, text, size, color, font=SANS, bold=False, italic=False,
         align=PP_ALIGN.LEFT, space_before=0, space_after=0, line_spacing=None,
         first=False):
    """Adiciona um paragrafo simples ao text frame."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line_spacing is not None:
        p.line_spacing = line_spacing
    style_run(p.add_run(), text, size, color, font=font, bold=bold, italic=italic)
    return p


def rich_para(tf, pieces, size, font=SANS, align=PP_ALIGN.LEFT, space_before=0,
              space_after=0, line_spacing=None, first=False):
    """Paragrafo com varios trechos: pieces = [(texto, cor, bold, font_opcional), ...]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line_spacing is not None:
        p.line_spacing = line_spacing
    for piece in pieces:
        txt, color = piece[0], piece[1]
        bold = piece[2] if len(piece) > 2 else False
        fnt = piece[3] if len(piece) > 3 else font
        style_run(p.add_run(), txt, size, color, font=fnt, bold=bold)
    return p


def body_block(slide, x, y, w, h, blocks, size=Pt(15), color=TEXT_DIM,
               line_spacing=1.35, space=10):
    """Paragrafos de texto corrido. Cada bloco pode ser str ou lista de trechos."""
    tf = textbox(slide, x, y, w, h)
    for i, block in enumerate(blocks):
        if isinstance(block, str):
            para(tf, block, size, color, line_spacing=line_spacing,
                 space_before=0 if i == 0 else space, first=(i == 0))
        else:
            rich_para(tf, block, size, line_spacing=line_spacing,
                      space_before=0 if i == 0 else space, first=(i == 0))
    return tf


# --------------------------------------------------- REALCE DE SINTAXE ------

KEYWORDS = {
    "import", "from", "function", "const", "let", "var", "return", "if", "else",
    "export", "default", "new", "true", "false", "null", "undefined", "typeof",
}
KNOWN_FUNCS = {
    "useState", "setEmail", "setSenha", "setErro", "setVerSenha", "alert",
    "preventDefault", "handleSubmit", "Login", "App", "console",
    "useNavigate", "navigate", "Home", "BrowserRouter", "Routes", "Route",
    "Link", "createRoot", "getElementById", "render",
    "Navigate", "RotaPrivada", "NaoEncontrada", "localStorage", "setItem",
    "getItem", "removeItem", "sair", "setLogado",
}

JS_RE = re.compile(
    r"""(?P<comment>//[^\n]*)
      | (?P<string>"[^"]*"|'[^']*')
      | (?P<arrow>=>)
      | (?P<tag></?[A-Za-z][\w.]*|/>)
      | (?P<num>\b\d+(?:\.\d+)?\b)
      | (?P<ident>[A-Za-z_$][\w$]*)
      | (?P<ws>\s+)
      | (?P<punct>.)""",
    re.VERBOSE,
)


def tokenize_js(line):
    """Devolve [(texto, cor), ...] para uma linha de JS/JSX."""
    if line.strip() in ("...", "…"):
        return [(line, C_COMMENT)]
    out = []
    for m in JS_RE.finditer(line):
        kind = m.lastgroup
        txt = m.group()
        if kind == "comment":
            color = C_COMMENT
        elif kind == "string":
            color = C_STRING
        elif kind == "tag":
            color = C_TAG
        elif kind == "num":
            color = C_NUM
        elif kind == "ident":
            if txt in KEYWORDS:
                color = C_KEYWORD
            elif txt in KNOWN_FUNCS or line[m.end():m.end() + 1] == "(":
                color = C_FUNC
            else:
                color = C_VAR
        elif kind == "ws":
            color = C_PUNCT
        else:  # arrow, punct
            color = C_PUNCT
        out.append((txt, color))
    return out or [(line, C_PUNCT)]


CSS_VALUE_RE = re.compile(r"(#[0-9a-fA-F]{3,8}|\b\d+(?:\.\d+)?\w*\b|[^,\s]+)")


def tokenize_css(line):
    """Devolve [(texto, cor), ...] para uma linha de CSS."""
    stripped = line.strip()
    if not stripped:
        return [(line, C_PUNCT)]
    # tambem pega a continuacao de um comentario de duas linhas
    if stripped.startswith("/*") or stripped.endswith("*/"):
        return [(line, C_COMMENT)]
    if stripped == "}":
        return [(line, C_PUNCT)]
    if stripped.endswith("{"):
        head = line[: line.index("{")]
        return [(head, C_SELECTOR), ("{", C_PUNCT)]
    if ":" in line:
        prop, value = line.split(":", 1)
        out = [(prop, C_VAR), (":", C_PUNCT)]
        pos = 0
        for m in CSS_VALUE_RE.finditer(value):
            if m.start() > pos:
                out.append((value[pos:m.start()], C_PUNCT))
            token = m.group()
            if token.startswith("#") or re.match(r"^\d", token):
                out.append((token, C_NUM))
            else:
                out.append((token, C_STRING))
            pos = m.end()
        if pos < len(value):
            out.append((value[pos:], C_PUNCT))
        return out
    return [(line, C_PUNCT)]


PY_KEYWORDS = {
    "from", "import", "class", "def", "return", "if", "elif", "else", "and",
    "or", "not", "raise", "True", "False", "None", "in", "is", "as", "with",
}
# nomes de tipo e de classe: pintados como tag, igual a um componente no JSX
PY_TIPOS = {
    "FastAPI", "HTTPException", "BaseModel", "DadosLogin", "str", "int",
    "dict", "bool",
}

PY_RE = re.compile(
    r"""(?P<comment>\#[^\n]*)
      | (?P<string>"[^"]*"|'[^']*')
      | (?P<decorator>@[A-Za-z_][\w.]*)
      | (?P<num>\b\d+(?:\.\d+)?\b)
      | (?P<ident>[A-Za-z_][\w]*)
      | (?P<ws>\s+)
      | (?P<punct>.)""",
    re.VERBOSE,
)


def tokenize_py(line):
    """Devolve [(texto, cor), ...] para uma linha de Python."""
    if line.strip() in ("...", "…"):
        return [(line, C_COMMENT)]
    out = []
    for m in PY_RE.finditer(line):
        kind = m.lastgroup
        txt = m.group()
        if kind == "comment":
            color = C_COMMENT
        elif kind == "string":
            color = C_STRING
        elif kind == "decorator":
            color = C_FUNC
        elif kind == "num":
            color = C_NUM
        elif kind == "ident":
            if txt in PY_KEYWORDS:
                color = C_KEYWORD
            elif txt in PY_TIPOS:
                color = C_TAG
            elif line[m.end():m.end() + 1] == "(":
                color = C_FUNC
            else:
                color = C_VAR
        else:
            color = C_PUNCT
        out.append((txt, color))
    return out or [(line, C_PUNCT)]


PROMPT_RE = re.compile(r"^((?:\([\w-]+\) )?(?:PS [^>]*>|C:\[^>]*>|\$) ?)(.*)$")


def tokenize_shell(line):
    """Terminal: prompt apagado, comando em destaque, saida em texto comum.

    Linha que comeca com PS ...> (ou C:\...> do cmd, ou $ do Git Bash) e
    comando digitado pelo aluno; o resto e o que o terminal responde.
    """
    if not line.strip():
        return [(line, C_PUNCT)]
    if line.lstrip().startswith("#"):
        return [(line, C_COMMENT)]
    m = PROMPT_RE.match(line)
    if not m:
        return [(line, TEXT_DIM)]
    prompt, cmd = m.groups()
    out = []
    if prompt.startswith("("):
        fim = prompt.index(")") + 1
        out.append((prompt[:fim], OK))
        prompt = prompt[fim:]
    out.append((prompt, TEXT_FAINT))
    for i, pedaco in enumerate(re.findall(r'"[^"]*"|\S+|\s+', cmd)):
        if pedaco.isspace():
            out.append((pedaco, C_PUNCT))
        elif i == 0:
            out.append((pedaco, C_FUNC))
        elif pedaco.startswith('"'):
            out.append((pedaco, C_STRING))
        elif pedaco.startswith("-"):
            out.append((pedaco, C_KEYWORD))
        else:
            out.append((pedaco, C_VAR))
    return out


def tokenize_toml(line):
    """pyproject.toml: [secao], chave = valor."""
    t = line.strip()
    if not t:
        return [(line, C_PUNCT)]
    if t.startswith("#"):
        return [(line, C_COMMENT)]
    if t.startswith("["):
        return [(line, C_SELECTOR)]
    if "=" in line and not t.startswith('"'):
        chave, valor = line.split("=", 1)
        return [(chave, C_VAR), ("=", C_PUNCT)] + tokenize_py(valor)
    return tokenize_py(line)


def tokenize_arvore(line):
    """Arvore de pastas: nome em destaque, comentario # apagado."""
    nome, _, comentario = line.partition("#")
    out = []
    m = re.match(r"^([\s│├└─]*)(.*?)(\s*)$", nome)
    galhos, arq, sobra = m.groups()
    out.append((galhos, TEXT_FAINT))
    out.append((arq, C_SELECTOR if arq.endswith("/") else C_VAR))
    out.append((sobra, C_PUNCT))
    if comentario:
        out.append(("#" + comentario, C_COMMENT))
    return out


TOKENIZADORES = {
    "js": tokenize_js, "css": tokenize_css, "py": tokenize_py,
    "shell": tokenize_shell, "toml": tokenize_toml, "arvore": tokenize_arvore,
}


# ------------------------------------------------------ PAINEL DE CODIGO ----

def code_panel(slide, x, y, w, h, filename, lines, highlight=None, lang="js",
               size=CODE_SIZE, line_h=CODE_LH, start=1):
    """
    Painel escuro com barra de janela e o codigo dentro.
    highlight = (primeira_linha, ultima_linha) com indice 0, inclusive.
    As linhas fora do destaque ficam apagadas.
    start = numero da primeira linha no gutter. So nao e 1 quando o arquivo
    continua num segundo painel (ver slide_codigo_duplo): a numeracao precisa
    seguir de onde parou, senao a segunda coluna parece outro arquivo.
    """
    bar_h = 0.42
    lh_in = line_h.pt / 72.0
    pad_x = 0.32
    pad_y = 0.26
    gutter_w = 0.46   # coluna dos numeros de linha, como num editor

    rect(slide, x, y, w, h, fill=SURFACE, line=BORDER,
         radius=0.035, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # barra de janela
    rect(slide, x, y, w, bar_h, fill=SURFACE_2, radius=0.09,
         shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    rect(slide, x, y + bar_h - 0.02, w, 0.02, fill=BORDER)
    for i, dot in enumerate((ERR, RGBColor(0xE0, 0xB0, 0x60), OK)):
        oval(slide, x + 0.22 + i * 0.22, y + 0.155, 0.11, 0.11, fill=dot)
    tf = textbox(slide, x + 0.95, y + 0.09, w - 1.2, 0.26)
    para(tf, filename, Pt(11), TEXT_DIM, font=MONO, first=True)

    # faixa de destaque
    if highlight is not None:
        a, b = highlight
        top = y + bar_h + pad_y + a * lh_in
        alt = (b - a + 1) * lh_in + 0.09
        rect(slide, x + 0.02 + gutter_w, top - 0.045,
             w - 0.10 - gutter_w, alt, fill=HILITE)
        rect(slide, x + 0.02 + gutter_w, top - 0.045, 0.035, alt, fill=ACCENT)

    # coluna de numeros de linha
    rect(slide, x + 0.02, y + bar_h, gutter_w, h - bar_h - 0.03, fill=SURFACE_2)
    rect(slide, x + 0.02 + gutter_w, y + bar_h, 0.012, h - bar_h - 0.03,
         fill=BORDER)
    tfn = textbox(slide, x, y + bar_h + pad_y, gutter_w - 0.06, h - bar_h - pad_y)
    for i in range(len(lines)):
        foco = highlight is not None and highlight[0] <= i <= highlight[1]
        p = tfn.paragraphs[0] if i == 0 else tfn.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.line_spacing = line_h
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        style_run(p.add_run(), str(i + start), size,
                  ACCENT if foco else TEXT_FAINT, font=MONO, bold=foco)

    tokenizer = TOKENIZADORES.get(lang, tokenize_js)
    tf = textbox(slide, x + pad_x + gutter_w, y + bar_h + pad_y,
                 w - 2 * pad_x - gutter_w, h - bar_h - 2 * pad_y)
    for i, line in enumerate(lines):
        dim = highlight is not None and not (highlight[0] <= i <= highlight[1])
        if not line.strip():
            pieces = [(" ", C_PUNCT)]
        elif dim:
            pieces = [(line, TEXT_FAINT)]
        else:
            pieces = tokenizer(line)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_h
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for txt, color in pieces:
            style_run(p.add_run(), txt, size, color, font=MONO)
    return tf


# ------------------------------------------------------- CROMO DO SLIDE ----

def eyebrow_pill(slide, texto, y=0.58):
    """Etiqueta da secao dentro de uma pilula com borda accent."""
    largura = 0.30 + len(texto) * 0.082
    rect(slide, M, y, largura, 0.32, fill=HILITE, line=ACCENT_DEEP,
         radius=0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = textbox(slide, M, y, largura, 0.32, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, texto.upper(), Pt(10), ACCENT, bold=True,
         align=PP_ALIGN.CENTER, first=True)
    return y + 0.32


def selo_escola(slide, x, y, size, pad=0.06):
    """Selo institucional: placa clara com o logo centralizado dentro.

    Sem o arquivo do logo desenha so a placa — o deck continua gerando.
    """
    rect(slide, x, y, size, size, fill=BRANCO,
         radius=0.13, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    if LOGO.exists():
        lado = size - 2 * pad
        slide.shapes.add_picture(str(LOGO), Inches(x + pad), Inches(y + pad),
                                 Inches(lado), Inches(lado))
    else:
        print(f"AVISO: logo nao encontrado em {LOGO} — selo saiu vazio.")
    return x + size


def rule(slide, x, y, w_accent=1.15, w_total=3.1):
    """Risco de titulo: barra accent curta sobre um trilho fino."""
    rect(slide, x, y + 0.016, w_total, 0.018, fill=BORDER)
    rect(slide, x, y, w_accent, 0.05, fill=ACCENT)


def chrome(slide, n):
    """Rodape comum: identificacao, numero e regua de progresso."""
    y = H - 0.52
    _ = n
    tf = textbox(slide, M, y, 6.6, 0.26)
    rich_para(tf, [("Programação Web 2", TEXT_FAINT),
                   ("  ·  ", ACCENT),
                   (AULA, TEXT_FAINT),
                   ("  ·  ", ACCENT),
                   (ESCOLA, TEXT_FAINT)], Pt(10), first=True)

    tf = textbox(slide, W - M - 2.0, y, 2.0, 0.26)
    rich_para(tf, [(f"{n:02d}", TEXT_DIM, True), (f" / {TOTAL_SLIDES}", TEXT_FAINT)],
              Pt(10), align=PP_ALIGN.RIGHT, first=True)

    track_w = W - 2 * M
    rect(slide, M, H - 0.16, track_w, 0.022, fill=BORDER)
    rect(slide, M, H - 0.16, track_w * n / TOTAL_SLIDES, 0.022, fill=ACCENT)


def header(slide, eyebrow, title, title_mono=False):
    """Cabecalho padrao: pilula da secao + titulo grande com risco de accent."""
    eyebrow_pill(slide, eyebrow)
    tf = textbox(slide, M, 1.02, W - 2 * M, 0.72)
    para(tf, title, Pt(36) if title_mono else Pt(40), TEXT,
         font=MONO if title_mono else DISPLAY, bold=True, first=True)
    rule(slide, M, 1.72)


# ---------------------------------------------------------------- LAYOUTS ---

def slide_capa(prs, eyebrow, titulo_1, titulo_2, subtitle, author,
               selo="FASTAPI"):
    slide = add_slide(prs)

    # halo por tras do desenho, para dar profundidade
    cx, cy = 9.95, 3.85
    oval(slide, cx - 3.2, cy - 3.2, 6.4, 6.4, fill=BG_HI)

    # as chaves de um JSON abracando a unica rota da API
    for dx, chave in ((-2.75, "{"), (1.45, "}")):
        tf = textbox(slide, cx + dx, cy - 1.45, 1.3, 2.6,
                     anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        para(tf, chave, Pt(150), ACCENT_DEEP, font=MONO, bold=True,
             align=PP_ALIGN.CENTER, first=True)
    rect(slide, cx - 1.3, cy - 0.3, 2.6, 0.6, fill=HILITE, line=ACCENT,
         line_w=Pt(1.5), radius=0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = textbox(slide, cx - 1.3, cy - 0.3, 2.6, 0.6, anchor=MSO_ANCHOR.MIDDLE)
    rich_para(tf, [("POST", C_KEYWORD, True), (" /login", C_STRING, True)],
              Pt(20), font=MONO, align=PP_ALIGN.CENTER, first=True)

    selo_escola(slide, M, 0.62, 1.06)

    eyebrow_pill(slide, eyebrow, y=1.92)

    tf = textbox(slide, M, 2.46, 7.8, 1.9)
    para(tf, titulo_1, Pt(60), TEXT, font=DISPLAY, bold=True,
         line_spacing=1.0, first=True)
    para(tf, titulo_2, Pt(60), ACCENT, font=DISPLAY, bold=True, line_spacing=1.0)

    rule(slide, M, 4.62, w_accent=1.6, w_total=4.2)

    tf = textbox(slide, M, 4.92, 7.4, 0.4)
    para(tf, subtitle, Pt(18), TEXT_DIM, first=True)

    tf = textbox(slide, M, 5.44, 7.4, 0.34)
    para(tf, author, Pt(13), TEXT_FAINT, first=True)

    rect(slide, M, H - 1.02, 1.42, 0.34, fill=HILITE, line=OK,
         radius=0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = textbox(slide, M, H - 1.02, 1.42, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, selo, Pt(11), OK, font=MONO, bold=True,
         align=PP_ALIGN.CENTER, first=True)

    rect(slide, M, H - 0.16, (W - 2 * M) / TOTAL_SLIDES, 0.022, fill=ACCENT)
    return slide


def slide_cartoes(prs, n, eyebrow, title, cards):
    """Tres cartoes numerados lado a lado."""
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    gap = 0.34
    cw = (W - 2 * M - 2 * gap) / 3
    top = 2.06
    ch = 4.5

    for i, card in enumerate(cards):
        x = M + i * (cw + gap)
        rect(slide, x, top, cw, ch, fill=SURFACE, line=BORDER,
             radius=0.028, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, x, top, cw, 0.045, fill=ACCENT)

        tf = textbox(slide, x + 0.34, top + 0.34, cw - 0.68, 0.42)
        para(tf, f"{i + 1:02d}", Pt(26), ACCENT_DEEP, font=MONO, bold=True, first=True)

        tf = textbox(slide, x + 0.34, top + 0.82, cw - 0.68, 0.36)
        para(tf, card["titulo"], Pt(22), TEXT, font=DISPLAY, bold=True, first=True)

        tf = textbox(slide, x + 0.34, top + 1.28, cw - 0.68, 2.0)
        para(tf, card["texto"], Pt(13), TEXT_DIM, line_spacing=1.34, first=True)

        # a caixa de comando comeca sempre na mesma altura nos tres cartoes
        box_top = top + 2.95
        if card.get("comando"):
            cmds = card["comando"]
            box_h = 0.28 + 0.23 * len(cmds)
            rect(slide, x + 0.24, box_top, cw - 0.48, box_h, fill=BG, line=BORDER,
                 radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            tf = textbox(slide, x + 0.38, box_top + 0.12, cw - 0.70, box_h - 0.24)
            for j, cmd in enumerate(cmds):
                # linha comecando com espaco = continuacao do comando anterior
                prefix = ("  ", ACCENT) if cmd.startswith(" ") else ("$ ", ACCENT)
                rich_para(tf, [prefix, (cmd.strip(), OK)], Pt(10),
                          font=MONO, line_spacing=Pt(16), first=(j == 0))
            box_top += box_h

        if card.get("nota"):
            tf = textbox(slide, x + 0.34, box_top + 0.16, cw - 0.68, 0.5)
            para(tf, card["nota"], Pt(10.5), TEXT_FAINT, italic=True,
                 line_spacing=1.2, first=True)

    chrome(slide, n)
    return slide


LABELS_JSX = [("JavaScript", ".js", C_KEYWORD), ("HTML", "<tags>", C_TAG),
              ("JSX", ".jsx", ACCENT)]


def slide_conceito(prs, n, eyebrow, title, destaque, blocos, labels=None):
    """Um conceito, muito respiro.

    A ilustracao da direita e uma soma de tres termos (A + B = C); o ultimo
    vem realcado. Passe `labels` como [(nome, detalhe, cor), ...].
    """
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    rect(slide, M, 2.25, 6.05, 1.2, fill=SURFACE, line=BORDER,
         radius=0.05, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, M, 2.25, 0.055, 1.2, fill=ACCENT)
    tf = textbox(slide, M + 0.42, 2.25, 5.35, 1.2, anchor=MSO_ANCHOR.MIDDLE)
    rich_para(tf, destaque, Pt(19), line_spacing=1.25, first=True)

    body_block(slide, M, 3.85, 6.05, 2.55, blocos, size=Pt(15))

    # ilustracao: A + B = C
    bx = 7.55
    bw = 5.15
    rect(slide, bx, 2.25, bw, 4.2, fill=SURFACE, line=BORDER,
         radius=0.03, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    labels = labels or LABELS_JSX
    row_h = 1.0
    op_h = 0.32
    bloco_h = 3 * row_h + 2 * op_h
    y0 = 2.25 + (4.2 - bloco_h) / 2
    for i, (nome, ext, cor) in enumerate(labels):
        y = y0 + i * (row_h + op_h)
        if i == 2:
            rect(slide, bx + 0.5, y, bw - 1.0, row_h, fill=HILITE,
                 radius=0.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = textbox(slide, bx + 0.5, y + 0.12, bw - 1.0, 0.42)
        para(tf, nome, Pt(19), TEXT if i == 2 else TEXT_DIM, bold=(i == 2),
             align=PP_ALIGN.CENTER, first=True)
        tf = textbox(slide, bx + 0.5, y + 0.56, bw - 1.0, 0.32)
        para(tf, ext, Pt(13), cor, font=MONO, align=PP_ALIGN.CENTER, first=True)
        if i < 2:
            tf = textbox(slide, bx + 0.5, y + row_h, bw - 1.0, op_h,
                         anchor=MSO_ANCHOR.MIDDLE)
            para(tf, "+" if i == 0 else "=", Pt(17), ACCENT, font=MONO,
                 bold=True, align=PP_ALIGN.CENTER, first=True)

    chrome(slide, n)
    return slide


AREA_TOP = 1.78          # topo da area de conteudo dos slides de codigo
AREA_H = 4.78            # altura util ate o rodape


def slide_codigo(prs, n, eyebrow, termo, filename, lines, highlight, blocos,
                 lang="py", nota=None, panel_w=None):
    """Layout principal: codigo a esquerda, explicacao a direita.

    A altura do painel acompanha o numero de linhas e ele fica centralizado
    verticalmente na area de conteudo — nada de painel meio vazio.
    """
    slide = add_slide(prs)

    eyebrow_pill(slide, eyebrow)
    panel_w = panel_w or PAINEL_LARGO
    tf = textbox(slide, M, 1.02, W - 2 * M, 0.66)
    para(tf, termo, Pt(31), TEXT, font=MONO, bold=True, first=True)
    rule(slide, M, 1.72)

    code_h = min(AREA_H, 0.42 + 0.52 + len(lines) * CODE_LH_IN)
    code_y = AREA_TOP + (AREA_H - code_h) / 2
    code_panel(slide, M, code_y, panel_w, code_h, filename, lines,
               highlight=highlight, lang=lang)

    x = M + panel_w + 0.38
    w = W - M - x
    rect(slide, x, AREA_TOP, 0.045, AREA_H, fill=BORDER)
    body_block(slide, x + 0.34, AREA_TOP + 0.08, w - 0.34, AREA_H - 1.4,
               blocos, size=Pt(14.5))

    if nota:
        ny = AREA_TOP + AREA_H - 1.18
        rect(slide, x + 0.34, ny, w - 0.34, 1.18, fill=SURFACE, line=BORDER,
             radius=0.05, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = textbox(slide, x + 0.58, ny + 0.1, w - 0.82, 0.98,
                     anchor=MSO_ANCHOR.MIDDLE)
        rich_para(tf, nota, Pt(12.5), line_spacing=1.25, first=True)

    chrome(slide, n)
    return slide


def slide_codigo_duplo(prs, n, eyebrow, termo, filename, col_a, col_b,
                       lang="py", legenda=None, larg_a=None,
                       size=DUPLO_SIZE, line_h=DUPLO_LH):
    """Um arquivo longo em duas colunas — a da direita CONTINUA a da esquerda.

    A numeracao do gutter segue corrida (a coluna B comeca onde a A parou) e a
    janela da direita se anuncia como continuacao. Sem isso as duas colunas
    numeram a partir de 1 e parecem dois arquivos diferentes, cada um cortado.
    """
    slide = add_slide(prs)

    eyebrow_pill(slide, eyebrow)
    tf = textbox(slide, M, 1.02, 8.0, 0.62)
    para(tf, termo, Pt(31), TEXT, font=MONO, bold=True, first=True)

    gap = 0.3
    wa = larg_a or (W - 2 * M - gap) / 2
    wb = W - 2 * M - gap - wa
    code_panel(slide, M, 1.62, wa, 5.05, f"{filename}   (1)", col_a, lang=lang,
               size=size, line_h=line_h)
    code_panel(slide, M + wa + gap, 1.62, wb, 5.05,
               f"{filename}   (2)  continuação", col_b, lang=lang,
               size=size, line_h=line_h, start=len(col_a) + 1)

    if legenda:
        tf = textbox(slide, M, 6.72, W - 2 * M, 0.3)
        rich_para(tf, legenda, Pt(12), first=True)

    chrome(slide, n)
    return slide


def slide_passos(prs, n, eyebrow, title, intro, passos, marcadores=None,
                 destaque_ultimo=False):
    """Passos numerados (A/B/C ou 01/02/03), um por linha."""
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    y = 1.95
    if intro:
        tf = textbox(slide, M, y, 9.6, 1.05)
        para(tf, intro, Pt(15), TEXT_DIM, line_spacing=1.35, first=True)
        y += 1.20

    marcadores = marcadores or [f"{i + 1:02d}" for i in range(len(passos))]
    row_h = 0.98
    gap = 0.22
    # centraliza o bloco de passos no espaco que sobrou
    bloco_h = len(passos) * (row_h + gap) - gap
    y = y + max(0.0, (6.72 - y - bloco_h) / 2)
    for i, passo in enumerate(passos):
        top = y + i * (row_h + gap)
        realce = destaque_ultimo and i == len(passos) - 1
        rect(slide, M, top, W - 2 * M, row_h,
             fill=HILITE if realce else SURFACE,
             line=ACCENT if realce else BORDER,
             radius=0.09, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

        tf = textbox(slide, M + 0.36, top, 0.9, row_h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, marcadores[i], Pt(24), ACCENT, font=MONO, bold=True, first=True)

        tw = W - 2 * M - 1.5
        tf = textbox(slide, M + 1.36, top, tw, row_h, anchor=MSO_ANCHOR.MIDDLE)
        if isinstance(passo, str):
            para(tf, passo, Pt(16), TEXT, line_spacing=1.25, first=True)
        else:
            rich_para(tf, passo, Pt(16), line_spacing=1.25, first=True)

    chrome(slide, n)
    return slide


# ------------------------------------------------- LAYOUTS DA AULA 06 ------
#
# A aula 6 sai do editor e vai para o terminal e para o navegador. Estes
# layouts cobrem o que a aula 5 nao tinha: a janela de terminal com a saida
# esperada, o print do Swagger, a tabela npm x uv, o desenho de ida e volta
# entre cliente e servidor e dois paineis lado a lado para comparar.

def slide_terminal(prs, n, eyebrow, title, linhas, blocos, nota=None,
                   janela="Terminal  ·  PowerShell"):
    """Terminal na largura toda: o comando e a saida que o aluno deve ver.

    Embaixo, a explicacao; se houver `nota`, ela vai numa caixa a direita.
    """
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    lh = TERM_LH.pt / 72.0
    top = 2.02
    code_h = 0.42 + 0.46 + len(linhas) * lh
    code_panel(slide, M, top, W - 2 * M, code_h, janela, linhas,
               lang="shell", size=TERM_SIZE, line_h=TERM_LH)

    y = top + code_h + 0.28
    fim = 6.78
    if nota:
        nw = 4.6
        body_block(slide, M, y, W - 2 * M - nw - 0.35, fim - y, blocos,
                   size=Pt(14), space=6, line_spacing=1.28)
        nx = W - M - nw
        nh = min(fim - y, 1.55)
        rect(slide, nx, y, nw, nh, fill=SURFACE, line=BORDER,
             radius=0.05, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, nx, y, 0.05, nh, fill=ACCENT)
        tf = textbox(slide, nx + 0.26, y + 0.12, nw - 0.44, nh - 0.24,
                     anchor=MSO_ANCHOR.MIDDLE)
        rich_para(tf, nota, Pt(12.5), line_spacing=1.22, first=True)
    else:
        body_block(slide, M, y, W - 2 * M, fim - y, blocos, size=Pt(14),
                   space=6, line_spacing=1.28)

    chrome(slide, n)
    return slide


def janela_navegador(slide, x, y, w, h_img, imagem, url):
    """Moldura de navegador (barra com endereco) com o print dentro."""
    bar_h = 0.42
    rect(slide, x - 0.02, y - 0.02, w + 0.04, bar_h + h_img + 0.04,
         fill=SURFACE_2, line=BORDER, radius=0.03,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    for i, dot in enumerate((ERR, RGBColor(0xE0, 0xB0, 0x60), OK)):
        oval(slide, x + 0.22 + i * 0.22, y + 0.155, 0.11, 0.11, fill=dot)
    rect(slide, x + 1.0, y + 0.08, w - 1.22, 0.26, fill=BG, line=BORDER,
         radius=0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = textbox(slide, x + 1.16, y + 0.08, w - 1.5, 0.26,
                 anchor=MSO_ANCHOR.MIDDLE)
    para(tf, url, Pt(10.5), TEXT_DIM, font=MONO, first=True)
    slide.shapes.add_picture(str(imagem), Inches(x), Inches(y + bar_h),
                             Inches(w), Inches(h_img))


def slide_imagem(prs, n, eyebrow, title, imagem, blocos, nota=None,
                 url="127.0.0.1:8000/docs", larg=8.2, alt_max=4.35,
                 topo=False):
    """Print real do navegador a esquerda, explicacao a direita.

    topo=True: print na largura toda e explicacao embaixo — para os prints
    baixos e largos (as respostas), que ficariam ilegiveis numa coluna.
    """
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    imagem = ASSETS / imagem
    px_w, px_h = PptxImage.from_file(str(imagem)).size
    if topo:
        w = W - 2 * M
        h = w * px_h / px_w
        janela_navegador(slide, M, 2.05, w, h, imagem, url)
        y = 2.05 + 0.42 + h + 0.35
        fim = 6.78
        nw = 4.6
        body_block(slide, M, y, W - 2 * M - (nw + 0.35 if nota else 0),
                   fim - y, blocos, size=Pt(14.5), space=6, line_spacing=1.3)
        if nota:
            nx = W - M - nw
            nh = min(fim - y, 1.55)
            rect(slide, nx, y, nw, nh, fill=SURFACE, line=BORDER,
                 radius=0.05, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            rect(slide, nx, y, 0.05, nh, fill=ACCENT)
            tf = textbox(slide, nx + 0.26, y + 0.12, nw - 0.44, nh - 0.24,
                         anchor=MSO_ANCHOR.MIDDLE)
            rich_para(tf, nota, Pt(12.5), line_spacing=1.22, first=True)
        chrome(slide, n)
        return slide

    w = larg
    h = w * px_h / px_w
    if h > alt_max:           # print alto demais: limita pela altura
        h = alt_max
        w = h * px_w / px_h
    total = 0.42 + h
    y = AREA_TOP + 0.1 + (AREA_H - 0.1 - total) / 2
    janela_navegador(slide, M, y, w, h, imagem, url)

    x = M + larg + 0.4
    cw = W - M - x
    rect(slide, x, AREA_TOP, 0.045, AREA_H, fill=BORDER)
    body_block(slide, x + 0.3, AREA_TOP + 0.08, cw - 0.3,
               AREA_H - (1.3 if nota else 0.1), blocos, size=Pt(14),
               line_spacing=1.3)
    if nota:
        ny = AREA_TOP + AREA_H - 1.12
        rect(slide, x + 0.3, ny, cw - 0.3, 1.12, fill=SURFACE, line=BORDER,
             radius=0.05, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = textbox(slide, x + 0.5, ny + 0.12, cw - 0.7, 0.9,
                     anchor=MSO_ANCHOR.MIDDLE)
        rich_para(tf, nota, Pt(12), line_spacing=1.22, first=True)

    chrome(slide, n)
    return slide


def celula(slide, x, y, w, h, conteudo, size, cor=TEXT_DIM, bold=False):
    """Texto de uma celula de tabela: str simples ou lista de trechos."""
    tf = textbox(slide, x + 0.22, y, w - 0.3, h, anchor=MSO_ANCHOR.MIDDLE)
    if isinstance(conteudo, str):
        para(tf, conteudo, size, cor, bold=bold, first=True)
    else:
        rich_para(tf, conteudo, size, first=True)


def slide_tabela(prs, n, eyebrow, title, cabecalho, linhas, larguras,
                 rodape=None):
    """Tabela de comparacao; a primeira coluna e o nome da coisa."""
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    top = 2.08
    head_h = 0.5
    row_h = 0.5
    total_w = W - 2 * M
    xs = [M]
    for frac in larguras[:-1]:
        xs.append(xs[-1] + frac * total_w)
    ws = [frac * total_w for frac in larguras]

    rect(slide, M, top, total_w, head_h, fill=HILITE, line=ACCENT_DEEP,
         radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    for x, w, txt in zip(xs, ws, cabecalho):
        celula(slide, x, top, w, head_h, txt.upper(), Pt(11), ACCENT, bold=True)

    for i, linha in enumerate(linhas):
        y = top + head_h + 0.06 + i * row_h
        if i % 2 == 0:
            rect(slide, M, y, total_w, row_h, fill=SURFACE)
        for j, (x, w, conteudo) in enumerate(zip(xs, ws, linha)):
            celula(slide, x, y, w, row_h, conteudo, Pt(14),
                   TEXT if j == 0 else TEXT_DIM, bold=(j == 0))

    if rodape:
        y = top + head_h + 0.06 + len(linhas) * row_h + 0.22
        tf = textbox(slide, M, y, total_w, 6.8 - y)
        rich_para(tf, rodape, Pt(14), line_spacing=1.3, first=True)

    chrome(slide, n)
    return slide


def seta(slide, x0, x1, y, para_direita=True, cor=ACCENT):
    """Seta horizontal fina com ponta triangular."""
    ponta = 0.2
    if para_direita:
        rect(slide, x0, y - 0.015, x1 - x0 - ponta + 0.02, 0.03, fill=cor)
        tri = rect(slide, x1 - ponta, y - ponta / 2, ponta, ponta, fill=cor,
                   shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        tri.rotation = 90
    else:
        rect(slide, x0 + ponta - 0.02, y - 0.015, x1 - x0 - ponta + 0.02, 0.03,
             fill=cor)
        tri = rect(slide, x0, y - ponta / 2, ponta, ponta, fill=cor,
                   shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        tri.rotation = 270


def slide_fluxo(prs, n, eyebrow, title, cliente, servidor, ida, volta, blocos):
    """Cliente a esquerda, servidor a direita, pedido indo e resposta voltando.

    cliente/servidor = {"nome", "detalhe", "texto"}
    ida/volta        = {"rotulo", "linhas": [trechos...]}
    """
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    top = 2.08
    bh = 2.72
    bw = 3.35
    for x, lado in ((M, cliente), (W - M - bw, servidor)):
        rect(slide, x, top, bw, bh, fill=SURFACE, line=BORDER,
             radius=0.04, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, x, top, bw, 0.045, fill=ACCENT)
        tf = textbox(slide, x + 0.3, top + 0.3, bw - 0.6, 0.5)
        para(tf, lado["nome"], Pt(26), TEXT, font=DISPLAY, bold=True,
             first=True)
        tf = textbox(slide, x + 0.3, top + 0.84, bw - 0.6, 0.34)
        para(tf, lado["detalhe"], Pt(13), C_TAG, font=MONO, first=True)
        tf = textbox(slide, x + 0.3, top + 1.3, bw - 0.6, bh - 1.4)
        para(tf, lado["texto"], Pt(13), TEXT_DIM, line_spacing=1.3, first=True)

    x0 = M + bw + 0.28
    x1 = W - M - bw - 0.28
    for k, (mov, direita) in enumerate(((ida, True), (volta, False))):
        y_seta = top + 0.42 + k * 1.36
        cor = OK if direita else C_SELECTOR
        tf = textbox(slide, x0, y_seta - 0.36, x1 - x0, 0.28)
        para(tf, mov["rotulo"].upper(), Pt(10), cor, bold=True,
             align=PP_ALIGN.CENTER, first=True)
        seta(slide, x0, x1, y_seta, para_direita=direita, cor=cor)
        caixa_h = 0.18 + 0.27 * len(mov["linhas"])
        rect(slide, x0 + 0.1, y_seta + 0.16, x1 - x0 - 0.2, caixa_h, fill=BG,
             line=BORDER, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = textbox(slide, x0 + 0.3, y_seta + 0.25, x1 - x0 - 0.5, caixa_h)
        for j, trechos in enumerate(mov["linhas"]):
            rich_para(tf, trechos, Pt(12), font=MONO, line_spacing=Pt(19),
                      first=(j == 0))

    body_block(slide, M, top + bh + 0.3, W - 2 * M, 6.8 - top - bh - 0.3,
               blocos, size=Pt(14.5), space=6, line_spacing=1.3)

    chrome(slide, n)
    return slide


def slide_comparar(prs, n, eyebrow, title, painel_a, painel_b, blocos,
                   larg_a=None):
    """Dois paineis lado a lado; painel = (titulo, linhas, lang)."""
    slide = add_slide(prs)
    header(slide, eyebrow, title)

    gap = 0.34
    wa = larg_a or (W - 2 * M - gap) / 2
    wb = W - 2 * M - gap - wa
    linhas = max(len(painel_a[1]), len(painel_b[1]))
    code_h = 0.42 + 0.52 + linhas * CODE_LH_IN
    top = 2.05
    code_panel(slide, M, top, wa, code_h, painel_a[0], painel_a[1],
               lang=painel_a[2])
    code_panel(slide, M + wa + gap, top, wb, code_h, painel_b[0], painel_b[1],
               lang=painel_b[2])

    y = top + code_h + 0.3
    body_block(slide, M, y, W - 2 * M, 6.8 - y, blocos, size=Pt(14.5),
               space=6, line_spacing=1.3)

    chrome(slide, n)
    return slide


# ------------------------------------------------- CONTEUDO DOS SLIDES ------
#
# Os blocos abaixo sao o codigo que aparece nos paineis. Regras de largura:
# no painel de codigo (PAINEL_LARGO) cabem ~55 colunas e ~14 linhas; no
# arquivo inteiro em duas colunas, 19 linhas por coluna. O Python nao deixa
# quebrar linha em qualquer lugar — por isso o main.py ja foi escrito com
# linhas curtas (email_certo / senha_certa, o dicionario aberto em linhas).
#
# "# ..." marca continuidade: o arquivo segue, mas nao interessa agora. E um
# comentario de proposito — o "..." solto, com recuo, vira IndentationError
# no aluno que digitar tudo o que ve.
# O conferir_codigo.py compara estes blocos com api-login/main.py e com o
# api-login/pyproject.toml.

CODE_TOPO = [
    "from fastapi import FastAPI, HTTPException",
    "from pydantic import BaseModel",
    "",
    "# a aplicacao: as rotas sao penduradas nela",
    'app = FastAPI(title="API do Ponto Eletrônico")',
    "",
    "# ...",
]

CODE_MODELO = [
    "# ...",
    "",
    "# o formato do corpo que o front-end manda",
    "class DadosLogin(BaseModel):",
    "    email: str",
    "    senha: str",
    "",
    "# ...",
]

CODE_ROTA = [
    "# ...",
    "",
    "# POST /login: confere o e-mail e a senha",
    '@app.post("/login")',
    "def login(dados: DadosLogin):",
    '    email_certo = dados.email == "user@email.com"',
    '    senha_certa = dados.senha == "123456"',
    "",
    "# ...",
]

CODE_IF = [
    "# ...",
    "def login(dados: DadosLogin):",
    '    email_certo = dados.email == "user@email.com"',
    '    senha_certa = dados.senha == "123456"',
    "",
    "    if email_certo and senha_certa:",
    "        # o dicionario vira JSON, status 200",
    "        return {",
    '            "mensagem": "Login realizado com sucesso!",',
    '            "email": dados.email,',
    "        }",
    "",
    "# ...",
]

CODE_RAISE = [
    "# ...",
    "    if email_certo and senha_certa:",
    "        # o dicionario vira JSON, status 200",
    "        return {",
    '            "mensagem": "Login realizado com sucesso!",',
    '            "email": dados.email,',
    "        }",
    "",
    "    # chegou aqui: algo nao bateu, status 401",
    "    raise HTTPException(",
    "        status_code=401,",
    '        detail="E-mail ou senha incorretos.",',
    "    )",
]

# o main.py inteiro, em duas colunas: B continua exatamente onde A parou
CODE_MAIN_A = [
    "from fastapi import FastAPI, HTTPException",
    "from pydantic import BaseModel",
    "",
    "# a aplicacao: as rotas sao penduradas nela",
    'app = FastAPI(title="API do Ponto Eletrônico")',
    "",
    "",
    "# o formato do corpo que o front-end manda",
    "class DadosLogin(BaseModel):",
    "    email: str",
    "    senha: str",
    "",
    "",
]

CODE_MAIN_B = [
    "# POST /login: confere o e-mail e a senha",
    '@app.post("/login")',
    "def login(dados: DadosLogin):",
    '    email_certo = dados.email == "user@email.com"',
    '    senha_certa = dados.senha == "123456"',
    "",
    "    if email_certo and senha_certa:",
    "        # o dicionario vira JSON, status 200",
    "        return {",
    '            "mensagem": "Login realizado com sucesso!",',
    '            "email": dados.email,',
    "        }",
    "",
    "    # chegou aqui: algo nao bateu, status 401",
    "    raise HTTPException(",
    "        status_code=401,",
    '        detail="E-mail ou senha incorretos.",',
    "    )",
]

TOML_PYPROJECT = [
    "[project]",
    'name = "api-login"',
    'version = "0.1.0"',
    'description = "Add your description here"',
    'readme = "README.md"',
    'requires-python = ">=3.12"',
    "dependencies = [",
    '    "fastapi[standard]==0.141.1",',
    "]",
]

ARVORE_INIT = [
    "api-login/",
    "├── .git/              # o repositório do Git",
    "├── .gitignore         # o que o Git não guarda",
    "├── .python-version    # a versão do Python: 3.12",
    "├── main.py            # o nosso código",
    "├── pyproject.toml     # a ficha do projeto",
    "└── README.md          # a descrição do projeto",
]

# so para o slide que ensina a ler os paineis — nao e para digitar
EXEMPLO_LEITURA = [
    "# ...",
    "def login(dados: DadosLogin):",
    '    email_certo = dados.email == "user@email.com"',
    '    senha_certa = dados.senha == "123456"',
    "",
    "    if email_certo and senha_certa:",
    "        # o dicionario vira JSON, status 200",
    "        return {",
    "# ...",
]

JSON_PEDIDO = [
    "{",
    '  "email": "user@email.com",',
    '  "senha": "123456"',
    "}",
]

PY_DICIONARIO = [
    "{",
    '    "mensagem": "Login realizado com sucesso!",',
    '    "email": dados.email,',
    "}",
]

# ---- o que o terminal mostra. Linha com PS ...> e o que o aluno digita. ----

TERM_INSTALAR_UV = [
    r'PS C:\projetos> powershell -ExecutionPolicy ByPass -c "irm https://astral.sh/uv/install.ps1 | iex"',
    r"Downloading uv (x86_64-pc-windows-msvc)",
    r"Installing to C:\Users\aluno\.local\bin",
    r"  uv.exe",
    r"  uvx.exe",
    r"everything's installed!",
]

TERM_VERSAO = [
    r"PS C:\projetos> uv --version",
    r"uv 0.6.10 (f2a2d982b 2025-03-25)",
]

TERM_INIT = [
    r"PS C:\projetos> uv init api-login --python 3.12",
    r"Initialized project `api-login` at `C:\projetos\api-login`",
    r"PS C:\projetos> cd api-login",
    r"PS C:\projetos\api-login>",
]

TERM_VENV = [
    r"PS C:\projetos\api-login> uv venv",
    r"Using CPython 3.12.9",
    r"Creating virtual environment at: .venv",
    r"Activate with: .venv\Scripts\activate",
]

TERM_ATIVAR = [
    r"PS C:\projetos\api-login> .venv\Scripts\activate",
    r"(api-login) PS C:\projetos\api-login>",
    r"",
    r"# no Prompt de Comando (cmd), o mesmo comando:",
    r"C:\projetos\api-login> .venv\Scripts\activate",
    r"# no Git Bash:",
    r"$ source .venv/Scripts/activate",
]

TERM_ADD = [
    r'(api-login) PS C:\projetos\api-login> uv add "fastapi[standard]==0.141.1"',
    r"Resolved 46 packages in 2.01s",
    r"Installed 44 packages in 5.23s",
    r" + fastapi==0.141.1",
    r" + pydantic==2.13.5",
    r" + uvicorn==0.53.0",
    r" ... e mais 41 pacotes",
]

TERM_DEV = [
    r"(api-login) PS C:\projetos\api-login> fastapi dev main.py",
    r"   Starting FastAPI in development mode",
    r"   Server started at http://127.0.0.1:8000",
    r"   Documentation at http://127.0.0.1:8000/docs",
    r"INFO:     Uvicorn running on http://127.0.0.1:8000 (Press CTRL+C to quit)",
    r"INFO:     Application startup complete.",
]

TERM_REABRIR = [
    r"# no dia seguinte: abrir a pasta, ativar e subir",
    r"PS C:\projetos> cd api-login",
    r"PS C:\projetos\api-login> .venv\Scripts\activate",
    r"(api-login) PS C:\projetos\api-login> fastapi dev main.py",
    r"",
    r"# projeto baixado do GitHub, sem a pasta .venv: o uv recria tudo",
    r"PS C:\projetos\api-login> uv sync",
]


def build(destino):
    """Monta o deck.

    Duas metades. A primeira e conceito puro: o aluno precisa saber O QUE e
    cada peca (API, HTTP, JSON, FastAPI, venv, uv) antes de digitar qualquer
    comando. A segunda e o guia de construcao em tres etapas — preparar o
    ambiente, escrever a API, rodar e testar no Swagger —, um comando por
    slide, sempre com a saida que o terminal deve mostrar.
    """
    prs = new_deck()

    # 01 - capa
    slide_capa(
        prs,
        eyebrow="Programação Web 2  ·  Aula 06",
        titulo_1="A API de",
        titulo_2="login",
        subtitle="Quem confere a senha agora é o servidor.",
        author=f"{PROFESSOR}  ·  {ESCOLA}",
    )

    # 02 - o mapa da aula
    slide_passos(
        prs, 2, "A API de login", "O que vamos fazer hoje", "",
        [
            [("Entender as peças: ", TEXT), ("cliente, servidor, API, HTTP e JSON", TEXT, True),
             (".", TEXT)],
            [("Entender as ferramentas: ", TEXT), ("FastAPI, venv e uv", TEXT, True),
             (" — e para que serve cada uma.", TEXT)],
            [("Construir uma API em Python com ", TEXT), ("uma rota só", TEXT, True),
             (": ", TEXT), ("POST /login", C_STRING, True, MONO), (".", TEXT)],
            [("Testar tudo no navegador, pelo ", TEXT), ("Swagger", TEXT, True),
             (" que o FastAPI gera sozinho.", TEXT)],
        ],
    )

    # ---------------------------------------------------------- CONCEITOS --

    # 03 - onde paramos
    slide_conceito(
        prs, 3, "Onde paramos", "O furo da aula 5",
        [("Quem confere a senha ", TEXT_DIM), ("não pode", TEXT, True),
         (" ser o navegador do visitante.", TEXT_DIM)],
        [[("Na aula passada, bastou digitar ", TEXT_DIM),
          ('localStorage.setItem("logado", "true")', C_FUNC, True, MONO),
          (" no console para entrar sem senha.", TEXT_DIM)],
         "E a senha certa está escrita no Login.jsx, que chega inteiro ao "
         "navegador: qualquer um lê.",
         "A saída: levar essa decisão para o servidor, onde o visitante não "
         "mexe."],
        labels=[("A senha", "escrita no Login.jsx", C_STRING),
                ("A sessão", "localStorage", C_VAR),
                ("Qualquer um", "lê e forja", ERR)],
    )

    # 04 - cliente e servidor
    slide_fluxo(
        prs, 4, "Conceito  ·  cliente e servidor", "Dois programas conversando",
        {"nome": "Cliente", "detalhe": "React  ·  navegador",
         "texto": "Desenha a tela e recolhe o que a pessoa digitou. Roda no "
                  "computador de quem usa — portanto não é confiável."},
        {"nome": "Servidor", "detalhe": "Python  ·  FastAPI",
         "texto": "Guarda as regras e os dados. Roda numa máquina nossa, que "
                  "o visitante não enxerga por dentro."},
        {"rotulo": "pedido", "linhas": [
            [("me deixa entrar?", TEXT_DIM)],
            [("e-mail + senha", C_STRING)]]},
        {"rotulo": "resposta", "linhas": [
            [("pode entrar", OK), ("  ou  ", TEXT_FAINT), ("não pode", ERR)]]},
        ["O cliente nunca decide sozinho: ele pergunta. Quem decide é o "
         "servidor, e a resposta dele é o que vale.",
         "Até hoje nosso sistema só tinha o lado esquerdo. Hoje nasce o direito."],
    )

    # 05 - o que e uma API
    slide_conceito(
        prs, 5, "Conceito  ·  API", "O que é uma API",
        [("API é o ", TEXT_DIM), ("cardápio", TEXT, True),
         (" do servidor: os pedidos que ele aceita.", TEXT_DIM)],
        ["API vem de Application Programming Interface: interface de "
         "programação. Não é uma tela para gente — é uma porta para outros "
         "programas.",
         [("Pense num restaurante: o cliente não entra na cozinha. Ele pede ao "
           "garçom um prato ", TEXT_DIM), ("que está no cardápio", TEXT, True),
          (", e o garçom traz o prato pronto.", TEXT_DIM)],
         "Nossa API vai ter um único item no cardápio: o login. Cada item "
         "desses se chama rota (ou endpoint)."],
        labels=[("O pedido", "POST /login", C_STRING),
                ("A API confere", "main.py", C_TAG),
                ("A resposta", "200 ou 401", ACCENT)],
    )

    # 06 - HTTP
    slide_fluxo(
        prs, 6, "Conceito  ·  HTTP", "Como é um pedido de verdade",
        {"nome": "Pedido", "detalhe": "request",
         "texto": "Método (o verbo), endereço (a rota) e, quando precisa "
                  "mandar dados, um corpo."},
        {"nome": "Resposta", "detalhe": "response",
         "texto": "Um código de status, que diz como foi, e um corpo com o "
                  "resultado."},
        {"rotulo": "o react manda", "linhas": [
            [("POST", C_KEYWORD, True), (" /login", C_STRING)],
            [('{"email": "…", "senha": "…"}', TEXT_DIM)]]},
        {"rotulo": "a api devolve", "linhas": [
            [("200", OK, True), (" OK", TEXT_DIM)],
            [('{"mensagem": "Login realizado…"}', TEXT_DIM)]]},
        [[("HTTP é o idioma dessa conversa — o mesmo que o navegador usa para "
           "abrir qualquer site. ", TEXT_DIM), ("GET", C_KEYWORD, True, MONO),
          (" busca algo; ", TEXT_DIM), ("POST", C_KEYWORD, True, MONO),
          (" envia dados para o servidor processar.", TEXT_DIM)],
         [("Login é POST: a senha vai no ", TEXT_DIM), ("corpo", TEXT, True),
          (" do pedido, e não pendurada no endereço, onde ficaria no "
           "histórico do navegador.", TEXT_DIM)]],
    )

    # 07 - codigos de status
    slide_passos(
        prs, 7, "Conceito  ·  HTTP", "O código de status diz como foi", "",
        [
            [("OK", TEXT, True), (" — deu certo. Todo pedido volta com um "
              "número desses. Aqui: e-mail e senha conferem.", TEXT)],
            [("Unauthorized", TEXT, True),
             (" — não autorizado. Aqui: e-mail ou senha errados.", TEXT)],
            [("Not Found", TEXT, True),
             (" — essa rota não existe. É o que aparece ao abrir a raiz /.", TEXT)],
            [("Unprocessable Entity", TEXT, True),
             (" — o corpo veio no formato errado, faltando um campo, por "
              "exemplo.", TEXT)],
        ],
        marcadores=["200", "401", "404", "422"],
    )

    # 08 - JSON
    slide_comparar(
        prs, 8, "Conceito  ·  JSON", "JSON: o formato da conversa",
        ("pedido  ·  o React manda (JSON)", JSON_PEDIDO, "js"),
        ("resposta  ·  o main.py devolve (dicionário)", PY_DICIONARIO, "py"),
        [[("JSON é texto com chaves e valores. O React fala JavaScript, a API "
           "fala Python — e os dois sabem ler e escrever ", TEXT_DIM),
          ("JSON", TEXT, True), (".", TEXT_DIM)],
         [("Em Python, a estrutura parecida se chama ", TEXT_DIM),
          ("dicionário", TEXT, True),
          (". O FastAPI converte sozinho: o dicionário que a rota devolve sai "
           "como JSON, e o JSON que chega vira um objeto Python.", TEXT_DIM)]],
        larg_a=4.9,
    )

    # 09 - FastAPI e quem vem junto
    slide_cartoes(
        prs, 9, "Conceito  ·  FastAPI", "FastAPI e quem vem junto",
        [
            {
                "titulo": "FastAPI",
                "texto": "O framework: escrevemos funções Python e ele as "
                         "transforma em rotas HTTP. O Pydantic, que vem junto, "
                         "confere o corpo do pedido.",
                "comando": ["@app.post(\"/login\")"],
                "nota": "Framework = o esqueleto pronto; a gente preenche.",
            },
            {
                "titulo": "Uvicorn",
                "texto": "O servidor: o programa que fica ligado escutando a "
                         "porta 8000 e entrega cada pedido que chega para o "
                         "FastAPI. Sem ele o código não atende ninguém.",
                "comando": ["fastapi dev main.py"],
                "nota": "O fastapi dev liga o Uvicorn por nós.",
            },
            {
                "titulo": "Swagger",
                "texto": "A documentação automática. O FastAPI lê as rotas e "
                         "monta uma página onde dá para ver e testar cada uma, "
                         "sem escrever uma linha de front-end.",
                "comando": ["http://127.0.0.1:8000/docs"],
                "nota": "Segue o padrão OpenAPI; /redoc mostra outra versão.",
            },
        ],
    )

    # 10 - o problema das dependencias
    slide_conceito(
        prs, 10, "Conceito  ·  ambiente", "O problema das dependências",
        [("Um Python só para todos os projetos ", TEXT_DIM),
         ("não dá certo", TEXT, True), (".", TEXT_DIM)],
        ["Dependência é um pacote que o projeto usa e que não veio com o "
         "Python: o FastAPI é uma. O instalador de pacotes padrão do Python "
         "chama pip.",
         "Instalando tudo no Python do computador, os projetos passam a "
         "dividir a mesma estante: o de hoje quer a versão 0.141, um antigo "
         "quer a 0.100 — e só cabe uma.",
         "Outro problema: o colega clona o projeto e não tem como saber quais "
         "pacotes (e quais versões) instalar."],
        labels=[("Projeto antigo", "fastapi 0.100", C_VAR),
                ("Projeto novo", "fastapi 0.141", C_VAR),
                ("Python do computador", "só cabe um", ERR)],
    )

    # 11 - venv
    slide_conceito(
        prs, 11, "Conceito  ·  venv", "venv: um Python só deste projeto",
        [("A pasta ", TEXT_DIM), (".venv", C_SELECTOR, True, MONO),
         (" é a estante particular do projeto.", TEXT_DIM)],
        [[("venv", TEXT, True), (" (virtual environment, ambiente virtual) é "
          "uma pasta dentro do projeto com um Python e os pacotes ", TEXT_DIM),
          ("só dele", TEXT, True), (". Cada projeto tem a sua; ninguém briga.", TEXT_DIM)],
         [("Ativar", TEXT, True), (" o venv é avisar ao terminal: “daqui para "
          "frente, python e fastapi são os desta pasta”. O sinal é o nome do "
          "projeto entre parênteses no começo da linha.", TEXT_DIM)],
         "A .venv não vai para o GitHub: é pesada e se recria em segundos."],
        labels=[("Python", "3.12", C_TAG),
                ("Pacotes", "fastapi, pydantic…", C_VAR),
                ("Ambiente do projeto", "api-login/.venv", ACCENT)],
    )

    # 12 - uv
    slide_cartoes(
        prs, 12, "Conceito  ·  uv", "uv: uma ferramenta para tudo isso",
        [
            {
                "titulo": "Instala o Python",
                "texto": "Pedimos a versão 3.12 e o uv resolve: usa a que já "
                         "existe na máquina ou baixa uma, sem mexer no Python "
                         "que você já tem.",
                "comando": ["uv init api-login --python 3.12"],
            },
            {
                "titulo": "Cria o venv",
                "texto": "Um comando cria a pasta .venv com o Python certo "
                         "dentro. É o mesmo venv de sempre do Python — o uv "
                         "só faz mais rápido.",
                "comando": ["uv venv"],
            },
            {
                "titulo": "Cuida dos pacotes",
                "texto": "Instala no venv, anota no pyproject.toml e grava as "
                         "versões exatas no uv.lock — assim o colega instala "
                         "igualzinho.",
                "comando": ["uv add \"fastapi[standard]\"", "uv sync"],
                "nota": "Faz o trabalho do pip, só que muito mais rápido.",
            },
        ],
    )

    # 13 - npm x uv
    slide_tabela(
        prs, 13, "Conceito  ·  uv", "Vocês já conhecem isso: é o npm do Python",
        ["O quê", "No React (npm)", "No Python (uv)"],
        [
            ["A ficha do projeto", [("package.json", C_VAR, False, MONO)],
             [("pyproject.toml", C_VAR, False, MONO)]],
            ["As versões exatas", [("package-lock.json", C_VAR, False, MONO)],
             [("uv.lock", C_VAR, False, MONO)]],
            ["Onde os pacotes moram", [("node_modules/", C_SELECTOR, False, MONO)],
             [(".venv/", C_SELECTOR, False, MONO)]],
            ["Adicionar um pacote", [("npm install react-router-dom", OK, False, MONO)],
             [('uv add "fastapi[standard]"', OK, False, MONO)]],
            ["Instalar o que a ficha pede", [("npm install", OK, False, MONO)],
             [("uv sync", OK, False, MONO)]],
            ["Rodar em desenvolvimento", [("npm run dev", OK, False, MONO)],
             [("fastapi dev main.py", OK, False, MONO)]],
            ["Endereço de desenvolvimento", [("localhost:5173", C_STRING, False, MONO)],
             [("127.0.0.1:8000", C_STRING, False, MONO)]],
        ],
        [0.30, 0.35, 0.35],
        rodape=[("A diferença: o node_modules isola sozinho; no Python o venv "
                 "precisa estar ", TEXT_DIM), ("ativo", TEXT, True),
                (".", TEXT_DIM)],
    )

    # ------------------------------------------------ ETAPA 1: AMBIENTE --

    # 14 - roteiro da etapa 1
    slide_passos(
        prs, 14, "Etapa 1 de 3", "Preparar o ambiente",
        "",
        [
            [("Crie a pasta ", TEXT), (r"C:\projetos", C_VAR, True, MONO),
             (" e abra no VS Code: ", TEXT), ("File → Open Folder", TEXT, True),
             (" (em português: Arquivo → Abrir Pasta).", TEXT)],
            [("Abra o terminal: ", TEXT), ("Terminal → New Terminal", TEXT, True),
             (" (Terminal → Novo Terminal). No canto do painel tem que aparecer ",
              TEXT), ("powershell", C_VAR, True, MONO), (".", TEXT)],
            [("Nos slides, digite ", TEXT), ("só o que vem depois do >", TEXT, True),
             (". O ", TEXT), (r"PS C:\projetos>", TEXT_DIM, True, MONO),
             (" quem escreve é o próprio terminal.", TEXT)],
            [("Um comando por slide: instalar o uv, criar o projeto e o ", TEXT),
             (".venv", C_SELECTOR, True, MONO),
             (", ativar e adicionar o FastAPI. Compare a saída antes de seguir.",
              TEXT)],
        ],
        marcadores=["A", "B", "C", "D"],
        destaque_ultimo=False,
    )

    # 15 - instalar o uv
    slide_terminal(
        prs, 15, "Etapa 1  ·  instalar o uv", "Instalando o uv",
        TERM_INSTALAR_UV,
        [[("Digite (ou copie) só a parte depois de ", TEXT_DIM),
          (r"PS C:\projetos>", TEXT, True, MONO),
          (" e aperte Enter. É uma linha só, e ", TEXT_DIM),
          ("só se faz uma vez", TEXT, True), (" por computador.", TEXT_DIM)],
         "O instalador coloca o uv numa pasta do usuário e a adiciona ao PATH "
         "— a lista de lugares onde o terminal procura programas.",
         "A versão e os detalhes da saída podem variar; o que importa é o "
         "everything's installed! no fim."],
        nota=[("Agora ", TEXT_DIM), ("feche o VS Code inteiro", TEXT, True),
              (" e abra de novo. O terminal só enxerga o PATH novo quando "
               "nasce depois da instalação.", TEXT_DIM)],
    )

    # 16 - conferir
    slide_terminal(
        prs, 16, "Etapa 1  ·  instalar o uv", "Conferindo a instalação",
        TERM_VERSAO,
        [[("Se aparecer uv seguido de um número de versão, está pronto. O "
           "número no seu computador provavelmente vai ser outro — ", TEXT_DIM),
          ("tudo bem", TEXT, True), (".", TEXT_DIM)],
         [("Não precisa instalar o Python antes: o uv cuida disso no próximo "
           "passo.", TEXT_DIM)]],
        nota=[("Deu ", TEXT_DIM),
              ("O termo 'uv' não é reconhecido…", ERR, True, MONO),
              ("? O terminal é anterior à instalação. Feche o VS Code "
               "inteiro, abra de novo e repita.", TEXT_DIM)],
    )

    # 17 - uv init
    slide_terminal(
        prs, 17, "Etapa 1  ·  o projeto", "Criando o projeto",
        TERM_INIT,
        [[("uv init", C_FUNC, True, MONO), (" cria a pasta ", TEXT_DIM),
          ("api-login", C_VAR, True, MONO),
          (" com o esqueleto de um projeto Python. O ", TEXT_DIM),
          ("--python 3.12", C_KEYWORD, True, MONO),
          (" fixa a versão do Python — todo mundo da turma usa a mesma.", TEXT_DIM)],
         [("Depois, ", TEXT_DIM), ("cd api-login", C_FUNC, True, MONO),
          (" entra na pasta. Todos os comandos seguintes são dados ", TEXT_DIM),
          ("de dentro dela", TEXT, True), (".", TEXT_DIM)]],
        nota=[("Confira o fim da linha do terminal: tem que terminar em ", TEXT_DIM),
              (r"\api-login>", C_VAR, True, MONO),
              (". Se não terminar, os próximos passos falham.", TEXT_DIM)],
    )

    # 18 - o que o uv init criou
    slide_codigo(
        prs, 18, "Etapa 1  ·  o projeto", "O que o uv init criou",
        "Explorer  ·  api-login", ARVORE_INIT, None,
        [[("pyproject.toml", C_VAR, True, MONO),
          (" é a ficha do projeto: nome, versão do Python e dependências — o "
           "package.json do Python.", TEXT_DIM)],
         [(".python-version", C_VAR, True, MONO),
          (" diz ao uv qual Python usar neste projeto.", TEXT_DIM)],
         [("main.py", C_VAR, True, MONO),
          (" veio com um “Hello” de exemplo. Vamos trocar tudo por nossa "
           "API.", TEXT_DIM)]],
        lang="arvore",
        nota=[(".git e .gitignore só aparecem se o Git estiver "
               "instalado.", TEXT_DIM)],
    )

    # 19 - uv venv
    slide_terminal(
        prs, 19, "Etapa 1  ·  o venv", "Criando o ambiente virtual",
        TERM_VENV,
        [[("uv venv", C_FUNC, True, MONO), (" cria a pasta ", TEXT_DIM),
          (".venv", C_SELECTOR, True, MONO),
          (" dentro do projeto, com o Python 3.12. Se a máquina não tiver essa "
           "versão, o uv baixa antes — aparece uma linha Downloading a mais.",
           TEXT_DIM)],
         "Repare na última linha: o próprio uv já ensina o comando do próximo "
         "slide."],
        nota=[("Por enquanto a .venv tem só o Python. Os pacotes entram daqui "
               "a dois slides, com o ", TEXT_DIM), ("uv add", C_FUNC, True, MONO),
              (".", TEXT_DIM)],
    )

    # 20 - ativar
    slide_terminal(
        prs, 20, "Etapa 1  ·  o venv", "Ativando o venv",
        TERM_ATIVAR,
        [[("Deu certo quando aparece ", TEXT_DIM), ("(api-login)", OK, True, MONO),
          (" no começo da linha. A partir daí, python e fastapi são os da "
           ".venv.", TEXT_DIM)],
         "A ativação vale só para aquele terminal: abriu outro, ative de novo."],
        nota=[("Erro ", TEXT_DIM),
              ("a execução de scripts foi desabilitada neste sistema", ERR, True, MONO),
              ("? Rode uma vez ", TEXT_DIM),
              ("Set-ExecutionPolicy -Scope CurrentUser RemoteSigned", C_FUNC, True, MONO),
              (", responda S e ative de novo.", TEXT_DIM)],
    )

    # 21 - uv add
    slide_terminal(
        prs, 21, "Etapa 1  ·  os pacotes", "Adicionando o FastAPI",
        TERM_ADD,
        [[("[standard]", C_KEYWORD, True, MONO),
          (" pede o FastAPI com os extras: o Uvicorn e o comando fastapi. ", TEXT_DIM),
          ("==0.141.1", C_KEYWORD, True, MONO),
          (" fixa a versão — a mesma para toda a turma.", TEXT_DIM)],
         [("As aspas são obrigatórias: sem elas o terminal pode entender os "
           "colchetes do jeito dele.", TEXT_DIM)]],
        nota=[("São ~44 pacotes porque o FastAPI também tem as dependências "
               "dele. O uv resolve a árvore inteira sozinho.", TEXT_DIM)],
    )

    # 22 - o que mudou
    slide_codigo(
        prs, 22, "Etapa 1  ·  os pacotes", "O que o uv add mudou",
        "pyproject.toml", TOML_PYPROJECT, (6, 8),
        ["Três coisas, e nenhuma delas a gente edita à mão:",
         [("pyproject.toml", C_VAR, True, MONO),
          (" ganhou o FastAPI na lista de dependências.", TEXT_DIM)],
         [("uv.lock", C_VAR, True, MONO),
          (" apareceu com a versão exata de cada um dos 44 pacotes.", TEXT_DIM)],
         [(".venv", C_SELECTOR, True, MONO),
          (" recebeu os pacotes instalados de fato.", TEXT_DIM)]],
        lang="toml",
        nota=[("Para instalar ou remover pacote, use sempre ", TEXT_DIM),
              ("uv add", C_FUNC, True, MONO), (" e ", TEXT_DIM),
              ("uv remove", C_FUNC, True, MONO), (".", TEXT_DIM)],
    )

    # --------------------------------------------------- ETAPA 2: CODIGO --

    # 23 - roteiro da etapa 2
    slide_passos(
        prs, 23, "Etapa 2 de 3", "Escrever a API",
        "Um arquivo só. No Explorer do VS Code, abra api-login → main.py, "
        "apague tudo o que o uv init deixou nele e vá digitando pedaço por "
        "pedaço. Salve sempre com Ctrl+S.",
        [
            [("Importar o FastAPI e criar a ", TEXT), ("app", C_VAR, True, MONO),
             (".", TEXT)],
            [("Descrever o corpo do pedido com uma classe ", TEXT),
             ("BaseModel", C_TAG, True, MONO), (".", TEXT)],
            [("Criar a rota ", TEXT), ("POST /login", C_STRING, True, MONO),
             (": responder 200 se bater e 401 se não bater.", TEXT)],
        ],
        marcadores=["A", "B", "C"],
    )

    # 24 - como ler os paineis de codigo
    slide_codigo(
        prs, 24, "Antes de digitar", "Como ler os painéis",
        "exemplo  ·  não digite este", EXEMPLO_LEITURA, (5, 7),
        [[("Faixa roxa", TEXT, True), (": é o que você digita ", TEXT_DIM),
          ("agora", TEXT, True), (".", TEXT_DIM)],
         [("Linhas apagadas", TEXT, True),
          (": já estão no seu arquivo, dos slides anteriores. Aparecem só para "
           "mostrar ", TEXT_DIM), ("onde", TEXT, True),
          (" escrever. Não digite de novo.", TEXT_DIM)],
         [("# ...", C_COMMENT, True, MONO),
          (": um pedaço do arquivo que não cabe no slide. Não digite.", TEXT_DIM)]],
        nota=[("Trecho digitado duas vezes: sem erro nenhum, a API aceita ",
               TEXT_DIM), ("qualquer", TEXT, True), (" senha.", TEXT_DIM)],
    )

    # 25 - imports e app
    slide_codigo(
        prs, 25, "Escrevendo a API  ·  main.py", "FastAPI()",
        "main.py", CODE_TOPO, (0, 4),
        [[("from … import", C_KEYWORD, True, MONO),
          (" é o import do Python: traz o FastAPI e o HTTPException do pacote "
           "fastapi, e o BaseModel do pydantic.", TEXT_DIM)],
         [("app", C_VAR, True, MONO),
          (" é a aplicação. Todas as rotas vão ser penduradas nela.", TEXT_DIM)],
         [("O ", TEXT_DIM), ("title", C_VAR, True, MONO),
          (" é o nome que aparece no topo do Swagger.", TEXT_DIM)]],
        nota=[("Linha que começa com ", TEXT_DIM), ("#", C_COMMENT, True, MONO),
              (" é comentário: o Python ignora.", TEXT_DIM)],
    )

    # 26 - o modelo
    slide_codigo(
        prs, 26, "Escrevendo a API  ·  main.py", "BaseModel",
        "main.py  ·  digite só a faixa destacada", CODE_MODELO, (2, 5),
        [[("DadosLogin", C_TAG, True, MONO),
          (" descreve o corpo que o pedido precisa ter: um email e uma senha, "
           "os dois texto (", TEXT_DIM), ("str", C_TAG, True, MONO), (").", TEXT_DIM)],
         [("Herdar de ", TEXT_DIM), ("BaseModel", C_TAG, True, MONO),
          (" liga o Pydantic: se faltar um campo, a API recusa com 422 antes "
           "de rodar o nosso código.", TEXT_DIM)],
         "E o Swagger usa essa classe para montar o exemplo de corpo."],
        nota=[("Em Python, a indentação faz parte do código: os campos ficam ",
               TEXT_DIM), ("4 espaços", TEXT, True), (" para dentro da classe.", TEXT_DIM)],
    )

    # 27 - a rota
    slide_codigo(
        prs, 27, "Escrevendo a API  ·  main.py", '@app.post("/login")',
        "main.py  ·  digite só a faixa destacada", CODE_ROTA, (2, 6),
        [[("O ", TEXT_DIM), ("@app.post", C_FUNC, True, MONO),
          (" logo acima da função é um decorador: ele pendura a função na app "
           "como a rota ", TEXT_DIM), ("POST /login", C_STRING, True, MONO),
          (".", TEXT_DIM)],
         [("dados: DadosLogin", C_VAR, True, MONO),
          (" diz ao FastAPI para ler o corpo do pedido e entregar já no "
           "formato da classe.", TEXT_DIM)],
         [("Cada ", TEXT_DIM), ("==", C_PUNCT, True, MONO),
          (" compara e guarda True ou False.", TEXT_DIM)]],
        nota=[("Não esqueça os ", TEXT_DIM), ("dois-pontos", TEXT, True),
              (" no fim da linha do def.", TEXT_DIM)],
    )

    # 28 - deu certo
    slide_codigo(
        prs, 28, "Escrevendo a API  ·  main.py", "return",
        "main.py  ·  digite só a faixa destacada", CODE_IF, (5, 10),
        [[("Se as duas comparações deram True, a função ", TEXT_DIM),
          ("devolve um dicionário", TEXT, True), (".", TEXT_DIM)],
         [("O FastAPI transforma o dicionário em JSON e responde com status ",
           TEXT_DIM), ("200", OK, True, MONO), (" — sem a gente escrever nada "
           "sobre isso.", TEXT_DIM)],
         "O return encerra a função: o que vem depois não roda."],
        nota=[("Compare com a aula 4: é a mesma conferência do handleSubmit — "
               "só mudou de lado.", TEXT_DIM)],
    )

    # 29 - deu errado
    slide_codigo(
        prs, 29, "Escrevendo a API  ·  main.py", "HTTPException",
        "main.py  ·  digite só a faixa destacada", CODE_RAISE, (8, 12),
        [[("Se o código chegou aqui, o if não entrou: algo não bateu.", TEXT_DIM)],
         [("raise HTTPException", C_KEYWORD, True, MONO),
          (" interrompe a função e responde na hora com o status ", TEXT_DIM),
          ("401", ERR, True, MONO), (" e o texto do detail.", TEXT_DIM)],
         "A mensagem não diz se o erro foi no e-mail ou na senha — de "
         "propósito: não damos pistas para quem tenta adivinhar."],
        nota=[("O raise fica no nível do if (4 espaços), ", TEXT_DIM),
              ("fora", TEXT, True), (" dele.", TEXT_DIM)],
    )

    # 30 - main.py inteiro
    slide_codigo_duplo(
        prs, 30, "Escrevendo a API", "main.py inteiro",
        "main.py", CODE_MAIN_A, CODE_MAIN_B, larg_a=5.3,
        legenda=[("Confira linha a linha, indentação inclusive. ", TEXT_DIM),
                 ("Nada a mais", TEXT, True),
                 (": um def login só. Depois, salve (Ctrl+S).", TEXT_DIM)],
    )

    # ---------------------------------------------------- ETAPA 3: TESTAR --

    # 31 - roteiro da etapa 3
    slide_passos(
        prs, 31, "Etapa 3 de 3", "Rodar e testar no Swagger",
        "O terminal fica ocupado com o servidor ligado. Não feche: se fechar, "
        "a API desliga.",
        [
            [("Subir o servidor com ", TEXT),
             ("fastapi dev main.py", C_FUNC, True, MONO), (".", TEXT)],
            [("Abrir ", TEXT), ("http://127.0.0.1:8000/docs", C_STRING, True, MONO),
             (" no navegador.", TEXT)],
            [("Testar três pedidos: senha certa, senha errada e corpo "
              "incompleto.", TEXT)],
        ],
        marcadores=["A", "B", "C"],
    )

    # 32 - fastapi dev
    slide_terminal(
        prs, 32, "Etapa 3  ·  rodar", "Ligando o servidor",
        TERM_DEV,
        [[("fastapi dev", C_FUNC, True, MONO),
          (" liga o Uvicorn apontando para a app do main.py. As linhas "
           "Server started e Documentation trazem os dois endereços.", TEXT_DIM)],
         [("Segure ", TEXT_DIM), ("Ctrl", TEXT, True),
          (" e clique no endereço do /docs para abrir no navegador. A saída "
           "real tem mais algumas linhas — tudo bem.", TEXT_DIM)],
         [("Modo dev", TEXT, True), (": salvou o main.py, o servidor reinicia "
          "sozinho. Para desligar, ", TEXT_DIM), ("Ctrl+C", TEXT, True),
          (" no terminal.", TEXT_DIM)]],
        nota=[("127.0.0.1", C_STRING, True, MONO),
              (" é o próprio computador (o mesmo que localhost) e ", TEXT_DIM),
              ("8000", C_NUM, True, MONO),
              (" é a porta. O Vite do React usava a 5173.", TEXT_DIM)],
    )

    # 33 - /docs
    slide_imagem(
        prs, 33, "Etapa 3  ·  Swagger", "http://127.0.0.1:8000/docs",
        "swagger-1-docs.png",
        [[("Esta página é o ", TEXT_DIM), ("Swagger", TEXT, True),
          (". Ninguém escreveu: o FastAPI montou lendo o main.py.", TEXT_DIM)],
         [("No topo, o ", TEXT_DIM), ("title", C_VAR, True, MONO),
          (" da app. Embaixo, a nossa única rota, com o verbo em verde.", TEXT_DIM)],
         [("Em Schemas está o ", TEXT_DIM), ("DadosLogin", C_TAG, True, MONO),
          (" — a classe virou documentação.", TEXT_DIM)]],
        nota=[("Clique na faixa verde ", TEXT_DIM),
              ("POST /login", C_STRING, True, MONO), (" para abrir a rota.", TEXT_DIM)],
    )

    # 34 - Try it out
    slide_imagem(
        prs, 34, "Etapa 3  ·  Swagger", "Try it out",
        "swagger-2-try-it-out.png",
        ["A rota aberta mostra o que ela espera: nenhum parâmetro no endereço "
         "e um Request body obrigatório (required).",
         [("O exemplo com ", TEXT_DIM), ('"string"', C_STRING, True, MONO),
          (" saiu da classe DadosLogin: dois campos de texto.", TEXT_DIM)],
         [("Clique em ", TEXT_DIM), ("Try it out", TEXT, True),
          (", à direita: o exemplo vira uma caixa editável.", TEXT_DIM)]],
    )

    # 35 - Execute
    slide_imagem(
        prs, 35, "Etapa 3  ·  Swagger", "Preenchendo e executando",
        "swagger-3-execute.png",
        [[("Troque os dois ", TEXT_DIM), ('"string"', C_STRING, True, MONO),
          (" pelas credenciais de teste:", TEXT_DIM)],
         [("user@email.com", C_STRING, True, MONO), (" e ", TEXT_DIM),
          ("123456", C_STRING, True, MONO), (".", TEXT_DIM)],
         [("Mantenha as aspas duplas e a vírgula entre os campos. Depois, ", TEXT_DIM),
          ("Execute", TEXT, True), (".", TEXT_DIM)]],
        nota=[("O Swagger faz o papel do React: monta o pedido POST e manda "
               "para a API de verdade.", TEXT_DIM)],
    )

    # 36 - 200
    slide_imagem(
        prs, 36, "Etapa 3  ·  testando", "Senha certa: 200",
        "swagger-4-200.png",
        [[("Role até ", TEXT_DIM), ("Server response", TEXT, True),
          (". Code ", TEXT_DIM), ("200", OK, True, MONO),
          (" e, no Response body, o dicionário do return virado JSON.", TEXT_DIM)],
         [("Ao mesmo tempo, o terminal registra o pedido numa linha nova: ",
           TEXT_DIM), ('"POST /login HTTP/1.1" 200 OK', C_STRING, True, MONO)]],
        nota=[("Acima do resultado, o Swagger mostra o mesmo pedido como "
               "comando curl — um jeito de chamar a API sem navegador.", TEXT_DIM)],
        topo=True,
    )

    # 37 - 401
    slide_imagem(
        prs, 37, "Etapa 3  ·  testando", "Senha errada: 401",
        "swagger-5-401.png",
        [[("Troque a senha por qualquer outra coisa e execute de novo. Agora "
           "a resposta é ", TEXT_DIM), ("401", ERR, True, MONO),
          (" com o detail do HTTPException: foi o raise que respondeu.", TEXT_DIM)],
         [("É ", TEXT_DIM), ("esta", TEXT, True),
          (" resposta que o Login.jsx vai receber para mostrar a mensagem de "
           "erro na tela.", TEXT_DIM)]],
        topo=True,
    )

    # 38 - 422
    slide_imagem(
        prs, 38, "Etapa 3  ·  testando", "Campo faltando: 422",
        "swagger-6-422.png",
        [[("Apague a linha da senha (e a vírgula que sobrou) e execute.", TEXT_DIM)],
         [("422", C_SELECTOR, True, MONO),
          (": nosso código nem rodou. O Pydantic viu que faltava ", TEXT_DIM),
          ("senha", C_VAR, True, MONO), (" no ", TEXT_DIM), ("body", C_VAR, True, MONO),
          (" e recusou — é o \"Field required\".", TEXT_DIM)],
         "Essa validação veio de graça, só por termos escrito a classe."],
        nota=[("Bônus: ", TEXT, True), ("http://127.0.0.1:8000/redoc", C_STRING, True, MONO),
              (" mostra a mesma documentação em outro formato, só de leitura.",
               TEXT_DIM)],
    )

    # ------------------------------------------------------- SEM ERROS ----

    # 39 - erros no ambiente
    slide_passos(
        prs, 39, "Quando der errado", "Erros no terminal", "",
        [
            [("O termo 'uv' não é reconhecido…", ERR, True, MONO),
             (" — feche o VS Code inteiro e abra de novo. Persistindo, "
              "reinstale o uv.", TEXT)],
            [("…Activate.ps1 não pode ser carregado…", ERR, True, MONO),
             (" — rode ", TEXT),
             ("Set-ExecutionPolicy -Scope CurrentUser RemoteSigned", C_FUNC, True, MONO),
             (" e responda S.", TEXT)],
            [("O termo 'fastapi' não é reconhecido…", ERR, True, MONO),
             (" — falta o (api-login) na linha: ative o venv, ou rode ", TEXT),
             ("uv run fastapi dev main.py", C_FUNC, True, MONO), (".", TEXT)],
            [("Path does not exist main.py", ERR, True, MONO),
             (" — você está fora da pasta. Rode ", TEXT),
             ("cd api-login", C_FUNC, True, MONO), (" e tente de novo.", TEXT)],
        ],
    )

    # 40 - erros ao testar
    slide_passos(
        prs, 40, "Quando der errado", "Erros ao testar", "",
        [
            [("O navegador não abre /docs", TEXT, True),
             (" — o terminal do fastapi dev foi fechado ou o Ctrl+C desligou "
              "o servidor. Suba de novo.", TEXT)],
            [('{"detail":"Not Found"}', ERR, True, MONO),
             (" em 127.0.0.1:8000 — é normal: a raiz / não é rota nossa. "
              "Acrescente /docs.", TEXT)],
            [("422 com a senha certa", TEXT, True),
             (" — o JSON está mal escrito: aspas simples, vírgula faltando "
              "entre os campos ou sobrando no último.", TEXT)],
            [("Salvei o main.py e nada mudou", TEXT, True),
             (" — aperte F5 na página do /docs: o Swagger não recarrega "
              "sozinho.", TEXT)],
        ],
    )

    # 41 - erros no main.py
    slide_passos(
        prs, 41, "Quando der errado", "Erros no main.py", "",
        [
            [("200 com ", TEXT), ("null", ERR, True, MONO),
             (" para qualquer senha — um trecho foi digitado duas vezes (dois "
              "def login). Compare com o slide 30.", TEXT)],
            [("IndentationError", ERR, True, MONO),
             (" no terminal — espaços errados no começo da linha. O terminal "
              "diz o número da linha.", TEXT)],
            [("SyntaxError: expected ':'", ERR, True, MONO),
             (" — faltou os dois-pontos no fim do ", TEXT),
             ("class", C_KEYWORD, True, MONO), (", do ", TEXT),
             ("def", C_KEYWORD, True, MONO), (" ou do ", TEXT),
             ("if", C_KEYWORD, True, MONO), (".", TEXT)],
            [("500 e ", TEXT),
             ("NameError: name 'HTTPException' is not defined", ERR, True, MONO),
             (" — faltou ele no import da primeira linha.", TEXT)],
        ],
    )

    # 42 - reabrindo
    slide_terminal(
        prs, 42, "No dia seguinte", "Voltando ao projeto",
        TERM_REABRIR,
        [[("Instalar o uv, ", TEXT_DIM), ("uv init", C_FUNC, True, MONO),
          (", ", TEXT_DIM), ("uv venv", C_FUNC, True, MONO), (" e ", TEXT_DIM),
          ("uv add", C_FUNC, True, MONO),
          (" foram uma vez só. Do segundo dia em diante são três comandos: "
           "entrar, ativar e subir.", TEXT_DIM)],
         [("Clonou o projeto em outro computador? A .venv não veio junto — ", TEXT_DIM),
          ("uv sync", C_FUNC, True, MONO),
          (" lê o uv.lock e recria o ambiente idêntico.", TEXT_DIM)]],
        nota=[("Porta 8000 ocupada por outro servidor esquecido? Ctrl+C no "
               "terminal dele, ou suba em outra: ", TEXT_DIM),
              ("fastapi dev main.py --port 8001", C_FUNC, True, MONO)],
    )

    # 43 - recapitulando
    slide_passos(
        prs, 43, "Recapitulando", "Do zero à API, em quatro passos", "",
        [
            [("Projeto e ambiente: ", TEXT),
             ("uv init api-login --python 3.12", C_FUNC, True, MONO),
             (" e ", TEXT), ("uv venv", C_FUNC, True, MONO), (".", TEXT)],
            [("Ativar com ", TEXT), (r".venv\Scripts\activate", C_FUNC, True, MONO),
             (" e adicionar o ", TEXT),
             ('uv add "fastapi[standard]==0.141.1"', C_FUNC, True, MONO), (".", TEXT)],
            [("No main.py: ", TEXT), ("app", C_VAR, True, MONO), (", ", TEXT),
             ("DadosLogin", C_TAG, True, MONO), (" e a rota ", TEXT),
             ('@app.post("/login")', C_FUNC, True, MONO), (".", TEXT)],
            [("Subir com ", TEXT), ("fastapi dev main.py", C_FUNC, True, MONO),
             (" e testar 200, 401 e 422 em ", TEXT), ("/docs", C_STRING, True, MONO),
             (".", TEXT)],
        ],
    )

    # 44 - o que ainda falta
    slide_conceito(
        prs, 44, "O que ainda falta", "Um primeiro passo, não o fim",
        [("A senha saiu do navegador — mas ", TEXT_DIM),
         ("ainda mora no código", TEXT, True), (".", TEXT_DIM)],
        ["O e-mail e a senha fixos no main.py são um exercício. Num sistema "
         "de verdade os usuários ficam num banco de dados, e a senha é "
         "guardada embaralhada (hash), nunca como texto.",
         "A API também ainda não entrega nada que prove o login para os "
         "próximos pedidos. Isso é o token, que vai substituir o \"true\" do "
         "localStorage.",
         "E o React ainda confere a senha sozinho: a próxima peça é o "
         "Login.jsx perguntando para esta API."],
        labels=[("Hoje", "POST /login", OK),
                ("Próximo", "o React chama a API", C_TAG),
                ("Depois", "banco + token", ACCENT)],
    )

    # 45 - trabalho
    slide_passos(
        prs, 45, "E agora?", "O teu trabalho", "",
        [
            [("Refazer a API ", TEXT), ("do zero", TEXT, True),
             (" em casa, numa pasta nova, seguindo só os slides.", TEXT)],
            [("Trocar o ", TEXT), ("title", C_VAR, True, MONO),
             (" da app pelo teu nome e conferir que ele aparece no Swagger.", TEXT)],
            [("Testar no /docs os três casos — ", TEXT), ("200", OK, True, MONO),
             (", ", TEXT), ("401", ERR, True, MONO), (" e ", TEXT),
             ("422", C_SELECTOR, True, MONO), (" — e tirar um print de cada.", TEXT)],
            [("Entregar os três prints e o ", TEXT), ("main.py", C_VAR, True, MONO),
             (". Sem a pasta .venv!", TEXT)],
        ],
        destaque_ultimo=True,
    )

    prs.save(destino)
    print(f"OK: {destino}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    saida = Path(__file__).with_name("Aula06-API-Login-FastAPI.pptx")
    build(str(saida))
