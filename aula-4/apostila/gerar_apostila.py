# -*- coding: utf-8 -*-
"""
Gera a apostila do trabalho da Aula 04 - Pagina de Login em React.

Segue o mesmo padrao da apostila da Aula 03 ("Instalacao do react.pdf"):
A4, capa, sumario, numeracao de pagina.

Uso:
    python gerar_apostila.py

Saida:
    Trabalho-Login-Aula04.docx

O sumario e criado como campo do Word: ao abrir o arquivo, clique com o botao
direito sobre ele e escolha "Atualizar campo" para preencher as paginas.
Depois e so exportar em PDF pelo proprio Word.
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

AUTOR = "Emanuel Nascente"
ESCOLA = "Escola Técnica Mesquita"
DISCIPLINA = "Programação Web 2"
DATA = "Aula 04"

INDIGO = RGBColor(0x4F, 0x46, 0xE5)
CINZA = RGBColor(0x55, 0x55, 0x60)
CINZA_CLARO = RGBColor(0x80, 0x80, 0x8C)
PRETO = RGBColor(0x1A, 0x1A, 0x1A)
VERDE = RGBColor(0x0A, 0x7A, 0x50)
VERMELHO = RGBColor(0xC0, 0x39, 0x3C)

MONO = "Consolas"
SANS = "Segoe UI"


# ----------------------------------------------------------------- BASE ----

def set_cell_bg(cell, hexcolor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(shd)


def add_field(paragraph, instrucao):
    """Insere um campo do Word (usado no sumario e na numeracao de pagina)."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instrucao
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, fld_end):
        run._r.append(el)
    return run


def atualizar_campos_ao_abrir(doc):
    """Faz o Word preencher o sumario sozinho na primeira vez que abrir."""
    settings = doc.settings.element
    tag = settings.find(qn("w:updateFields"))
    if tag is None:
        tag = OxmlElement("w:updateFields")
        settings.append(tag)
    tag.set(qn("w:val"), "true")


def estilo_base(doc):
    normal = doc.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(11)
    normal.font.color.rgb = PRETO
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    for nome, tamanho, cor in (("Heading 1", 20, INDIGO),
                               ("Heading 2", 14, PRETO),
                               ("Heading 3", 12, CINZA)):
        st = doc.styles[nome]
        st.font.name = SANS
        st.font.size = Pt(tamanho)
        st.font.bold = True
        st.font.color.rgb = cor
        st.paragraph_format.space_before = Pt(16 if nome == "Heading 1" else 12)
        st.paragraph_format.space_after = Pt(6)


def rodape_numerado(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{DISCIPLINA}  ·  {DATA}  ·  ")
    run.font.size = Pt(8)
    run.font.color.rgb = CINZA_CLARO
    run.font.name = SANS
    campo = add_field(p, "PAGE")
    campo.font.size = Pt(8)
    campo.font.color.rgb = CINZA_CLARO
    campo.font.name = SANS


# ------------------------------------------------------------- BLOCOS ------

def paragrafo(doc, texto, size=11, cor=PRETO, italico=False, negrito=False,
              align=None, space_after=8):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(texto)
    run.font.size = Pt(size)
    run.font.color.rgb = cor
    run.font.italic = italico
    run.font.bold = negrito
    run.font.name = SANS
    return p


def rico(doc, pedacos, size=11, align=None, space_after=8):
    """pedacos = [(texto, negrito, mono, cor), ...]"""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for pedaco in pedacos:
        if isinstance(pedaco, str):
            pedaco = (pedaco,)
        txt = pedaco[0]
        negrito = pedaco[1] if len(pedaco) > 1 else False
        mono = pedaco[2] if len(pedaco) > 2 else False
        cor = pedaco[3] if len(pedaco) > 3 else PRETO
        run = p.add_run(txt)
        run.font.size = Pt(size - 0.5 if mono else size)
        run.font.bold = negrito
        run.font.color.rgb = cor
        run.font.name = MONO if mono else SANS
    return p


def lista(doc, itens, marcador="•"):
    for item in itens:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"{marcador}  ")
        run.font.color.rgb = INDIGO
        run.font.bold = True
        run.font.name = SANS
        if isinstance(item, str):
            r = p.add_run(item)
            r.font.name = SANS
            r.font.size = Pt(11)
        else:
            for pedaco in item:
                if isinstance(pedaco, str):
                    pedaco = (pedaco,)
                txt = pedaco[0]
                negrito = pedaco[1] if len(pedaco) > 1 else False
                mono = pedaco[2] if len(pedaco) > 2 else False
                r = p.add_run(txt)
                r.font.bold = negrito
                r.font.name = MONO if mono else SANS
                r.font.size = Pt(10.5 if mono else 11)


def codigo(doc, linhas, legenda=None):
    """Bloco de codigo em tabela de uma celula com fundo cinza."""
    tabela = doc.add_table(rows=1, cols=1)
    tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = tabela.cell(0, 0)
    set_cell_bg(cell, "F4F4F7")
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for i, linha in enumerate(linhas):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(linha if linha else " ")
        run.font.name = MONO
        run.font.size = Pt(9.5)
        run.font.color.rgb = PRETO
    if legenda:
        paragrafo(doc, legenda, size=9, cor=CINZA_CLARO, italico=True,
                  space_after=12)
    else:
        paragrafo(doc, "", size=4, space_after=4)
    return tabela


def caixa(doc, titulo, texto, cor=INDIGO, fundo="EEF0FE"):
    """Caixa de destaque: dica, atencao, etc."""
    tabela = doc.add_table(rows=1, cols=1)
    cell = tabela.cell(0, 0)
    set_cell_bg(cell, fundo)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(titulo)
    run.font.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = cor
    run.font.name = SANS
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run = p2.add_run(texto)
    run.font.size = Pt(10.5)
    run.font.name = SANS
    paragrafo(doc, "", size=4, space_after=6)
    return tabela


def h1(doc, texto):
    return doc.add_heading(texto, level=1)


def h2(doc, texto):
    return doc.add_heading(texto, level=2)


def h3(doc, texto):
    return doc.add_heading(texto, level=3)


# --------------------------------------------------------------- CAPA ------

def capa(doc):
    for _ in range(4):
        paragrafo(doc, "", space_after=0)

    paragrafo(doc, DATA.upper(), size=11, cor=INDIGO, negrito=True, space_after=4)
    p = doc.add_paragraph()
    run = p.add_run("TRABALHO\nPÁGINA DE LOGIN")
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.name = SANS
    run.font.color.rgb = PRETO
    p.paragraph_format.space_after = Pt(10)

    paragrafo(doc, "React + Vite  ·  useState  ·  mostrar/ocultar senha",
              size=12, cor=CINZA, space_after=30)

    paragrafo(doc, AUTOR, size=12, negrito=True, space_after=2)
    paragrafo(doc, f"{DISCIPLINA}  ·  {ESCOLA}", size=11, cor=CINZA_CLARO)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def sumario(doc):
    h1(doc, "Sumário")
    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-2" \h \z \u')
    paragrafo(doc, "", size=9, space_after=4)
    paragrafo(doc,
              "O Word preenche este sumário ao abrir o arquivo. Se ele "
              "aparecer vazio, clique com o botão direito sobre ele e escolha "
              "\"Atualizar campo\".",
              size=9, cor=CINZA_CLARO, italico=True)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------ CONTEUDO -----

def conteudo(doc):
    # ---------------------------------------------------------- 1
    h1(doc, "1. O que você vai entregar")
    paragrafo(doc,
              "Na aula 4 nós construímos juntos a primeira página da nossa "
              "aplicação: uma tela de login em React, com estado, validação e "
              "estilo. Agora é a tua vez de levar essa tela adiante.")
    paragrafo(doc, "O trabalho tem três entregas obrigatórias:", negrito=True)
    lista(doc, [
        [("Botão mostrar / ocultar senha", True), (" — o campo de senha precisa "
          "poder alternar entre escondido e visível.")],
        [("CSS autoral", True), (" — a tela deve ter a tua identidade visual, "
          "e não a cópia exata do exemplo da aula.")],
        [("Segunda página iniciada", True), (" — o começo da próxima tela da "
          "tua aplicação, que vamos evoluir na aula 5.")],
    ])
    caixa(doc, "Antes de tudo",
          "Este trabalho parte do que fizemos em aula. Se o teu Login.jsx "
          "ainda não roda, comece pela seção 3 — ela reconstrói a base do "
          "zero.")

    # ---------------------------------------------------------- 2
    h1(doc, "2. Antes de começar")
    paragrafo(doc, "Confira se o teu ambiente está pronto. No terminal:")
    codigo(doc, [
        "node -v      # deve mostrar algo como v22.x",
        "npm -v       # deve mostrar algo como 10.x",
    ], "Se algum dos dois não responder, reinstale o Node.js "
       "(veja a apostila da aula 3).")
    paragrafo(doc, "E confira se você tem:")
    lista(doc, [
        "O projeto criado com o Vite, com a pasta node_modules já instalada.",
        [("O arquivo ", False), ("src/Login.jsx", False, True),
         (" com o componente da aula.", False)],
        [("O arquivo ", False), ("src/Login.css", False, True),
         (" com a estilização.", False)],
        [("O ", False), ("npm run dev", False, True),
         (" rodando e abrindo a tela no navegador.", False)],
    ])

    # ---------------------------------------------------------- 3
    h1(doc, "3. Etapa 1 — Recriando a base")
    paragrafo(doc,
              "Se você já tem o login da aula funcionando, pode pular direto "
              "para a Etapa 2. Senão, siga daqui.")

    h2(doc, "3.1 Criando o projeto")
    codigo(doc, [
        "npm create vite@latest meu-login -- --template react",
        "cd meu-login",
        "npm install",
        "npm run dev",
    ], "Atenção ao espaço antes dos dois traços: \"meu-login -- --template\".")

    h2(doc, "3.2 O componente Login")
    paragrafo(doc, "Crie o arquivo src/Login.jsx com este conteúdo:")
    codigo(doc, [
        'import { useState } from "react";',
        'import "./Login.css";',
        "",
        "function Login() {",
        '  const [email, setEmail] = useState("");',
        '  const [senha, setSenha] = useState("");',
        '  const [erro, setErro] = useState("");',
        "",
        "  function handleSubmit(e) {",
        "    e.preventDefault();",
        "",
        "    if (!email || !senha) {",
        '      setErro("Preencha todos os campos.");',
        "      return;",
        "    }",
        "",
        '    if (email === "user@email.com" && senha === "123456") {',
        '      alert("Login realizado com sucesso!");',
        '      setErro("");',
        "    } else {",
        '      setErro("E-mail ou senha incorretos.");',
        "    }",
        "  }",
        "",
        "  return (",
        '    <div className="container">',
        "      <form onSubmit={handleSubmit}>",
        "        <h2>Login</h2>",
        "",
        "        <input",
        '          type="email"',
        '          placeholder="Email"',
        "          value={email}",
        "          onChange={(e) => setEmail(e.target.value)}",
        "        />",
        "",
        "        <input",
        '          type="password"',
        '          placeholder="Senha"',
        "          value={senha}",
        "          onChange={(e) => setSenha(e.target.value)}",
        "        />",
        "",
        '        {erro && <p className="erro">{erro}</p>}',
        "",
        '        <button type="submit">Entrar</button>',
        "      </form>",
        "    </div>",
        "  );",
        "}",
        "",
        "export default Login;",
    ])

    paragrafo(doc, "E ligue o componente no src/App.jsx:")
    codigo(doc, [
        'import Login from "./Login";',
        "",
        "function App() {",
        "  return <Login />;",
        "}",
        "",
        "export default App;",
    ])
    caixa(doc, "Lembre-se",
          "Login e senha de teste: user@email.com / 123456. Essa validação "
          "está no front-end só para aprendermos o useState — em uma "
          "aplicação de verdade, quem confere login e senha é o servidor. "
          "É exatamente isso que faremos com o back-end em C# mais adiante.")

    # ---------------------------------------------------------- 4
    h1(doc, "4. Etapa 2 — Botão mostrar / ocultar senha")
    paragrafo(doc,
              "Essa é a parte central do trabalho. A ideia é simples: um botão "
              "que troca o campo de senha entre escondido e visível. Vamos "
              "fazer em quatro passos.")

    h2(doc, "4.1 Passo 1 — Criar o estado")
    paragrafo(doc,
              "Precisamos que o componente \"lembre\" se a senha está visível "
              "ou não. Isso é um estado — e como só existem duas situações "
              "(visível ou escondida), o valor é um booleano.")
    paragrafo(doc, "Junto com os outros useState, adicione:")
    codigo(doc, [
        "const [verSenha, setVerSenha] = useState(false);",
    ], "Começa em false porque, ao abrir a tela, a senha deve estar escondida.")

    h2(doc, "4.2 Passo 2 — Deixar o type dinâmico")
    paragrafo(doc,
              "Hoje o campo tem type=\"password\" fixo. Ele precisa virar "
              "\"text\" quando a senha estiver visível. Para isso usamos o "
              "operador ternário, que é um if/else escrito em uma linha só:")
    codigo(doc, [
        "condição ? valor_se_verdadeiro : valor_se_falso",
    ])
    paragrafo(doc, "Aplicando no nosso input:")
    codigo(doc, [
        'type={verSenha ? "text" : "password"}',
    ], "Lê-se: se verSenha for verdadeiro, use \"text\"; senão, use "
       "\"password\".")

    h2(doc, "4.3 Passo 3 — Criar o botão")
    paragrafo(doc,
              "Agora o botão que inverte o estado. Ele fica logo depois do "
              "input de senha:")
    codigo(doc, [
        "<button",
        '  type="button"',
        '  className="ver-senha"',
        "  onClick={() => setVerSenha(!verSenha)}",
        ">",
        '  {verSenha ? "Ocultar" : "Mostrar"}',
        "</button>",
    ])
    paragrafo(doc, "Três detalhes importantes nesse trecho:", negrito=True)
    lista(doc, [
        [("type=\"button\"", False, True), (" — sem isso o botão seria tratado "
          "como submit e enviaria o formulário a cada clique. Esse é o erro "
          "mais comum deste trabalho.", False)],
        [("onClick={() => setVerSenha(!verSenha)}", False, True),
         (" — o ponto de exclamação inverte o valor: se era false vira true, "
          "se era true vira false.", False)],
        [("{verSenha ? \"Ocultar\" : \"Mostrar\"}", False, True),
         (" — o texto do próprio botão também muda conforme o estado.", False)],
    ])

    h2(doc, "4.4 Passo 4 — Agrupar o campo e o botão")
    paragrafo(doc,
              "Para o botão ficar dentro do campo, envolvemos os dois em uma "
              "div e usamos position no CSS:")
    codigo(doc, [
        '<div className="campo-senha">',
        "  <input",
        '    type={verSenha ? "text" : "password"}',
        '    placeholder="Senha"',
        "    value={senha}",
        "    onChange={(e) => setSenha(e.target.value)}",
        "  />",
        "  <button",
        '    type="button"',
        '    className="ver-senha"',
        "    onClick={() => setVerSenha(!verSenha)}",
        "  >",
        '    {verSenha ? "Ocultar" : "Mostrar"}',
        "  </button>",
        "</div>",
    ])
    paragrafo(doc, "E no Login.css:")
    codigo(doc, [
        ".campo-senha {",
        "  position: relative;",
        "  display: flex;",
        "}",
        "",
        ".campo-senha input {",
        "  width: 100%;",
        "  padding-right: 5rem;",
        "}",
        "",
        ".ver-senha {",
        "  position: absolute;",
        "  right: 0.4rem;",
        "  top: 50%;",
        "  transform: translateY(-50%);",
        "  background: transparent;",
        "  color: #4f46e5;",
        "  border: none;",
        "  font-size: 0.8rem;",
        "  padding: 0.3rem 0.5rem;",
        "  cursor: pointer;",
        "  width: auto;",
        "}",
    ])
    caixa(doc, "Se der errado",
          "1) O formulário envia sozinho ao clicar no botão → faltou "
          "type=\"button\".  "
          "2) A senha não muda de visível para escondida → confira se o input "
          "usa type={verSenha ? ...} e não type=\"password\" fixo.  "
          "3) O botão aparece embaixo do campo em vez de dentro → faltou "
          "position: relative na .campo-senha.",
          cor=VERMELHO, fundo="FDEFEF")

    # ---------------------------------------------------------- 5
    h1(doc, "5. Etapa 3 — Estilizando do seu jeito")
    paragrafo(doc,
              "O CSS da aula é só um ponto de partida. A entrega pede uma tela "
              "com a tua cara — então mexa nele de verdade. Coisas que valem a "
              "pena mudar:")
    lista(doc, [
        [("Cores", True), (" — troque o #4f46e5 do botão e o #f0f2f5 do fundo "
          "por uma combinação sua.")],
        [("Cantos", True), (" — border-radius menor deixa a tela mais séria; "
          "maior, mais leve.")],
        [("Sombra", True), (" — box-shadow mais suave ou mais marcada muda "
          "bastante a sensação de profundidade.")],
        [("Hover", True), (" — dê um retorno visual quando o mouse passa "
          "sobre o botão.")],
        [("Tipografia", True), (" — font-family e font-size mudam o tom da "
          "página inteira.")],
    ])
    caixa(doc, "Cuidado com o contraste",
          "Texto claro sobre fundo claro (ou escuro sobre escuro) pode ficar "
          "bonito na tua tela e ilegível no projetor. Teste antes de "
          "entregar.")

    # ---------------------------------------------------------- 6
    h1(doc, "6. Etapa 4 — Começando a nova página")
    paragrafo(doc,
              "A terceira entrega é o começo da próxima tela. Ainda não vamos "
              "usar biblioteca de rotas — dá para trocar de página com o que "
              "já sabemos: um estado.")
    paragrafo(doc, "Crie o arquivo src/Home.jsx:")
    codigo(doc, [
        "function Home() {",
        "  return (",
        '    <div className="container">',
        "      <h2>Bem-vindo!</h2>",
        "      <p>Esta é a próxima página da nossa aplicação.</p>",
        "    </div>",
        "  );",
        "}",
        "",
        "export default Home;",
    ])
    paragrafo(doc, "E controle qual tela aparece a partir do App.jsx:")
    codigo(doc, [
        'import { useState } from "react";',
        'import Login from "./Login";',
        'import Home from "./Home";',
        "",
        "function App() {",
        "  const [logado, setLogado] = useState(false);",
        "",
        "  return logado ? <Home /> : <Login aoEntrar={() => setLogado(true)} />;",
        "}",
        "",
        "export default App;",
    ], "Para isso funcionar, o Login precisa chamar aoEntrar() no lugar do "
       "alert quando o login der certo.")
    caixa(doc, "Na aula 5",
          "Vamos transformar essa troca manual em navegação de verdade e "
          "conectar a tela ao back-end do nosso sistema de ponto eletrônico.")

    # ---------------------------------------------------------- 7
    h1(doc, "7. Como entregar")
    lista(doc, [
        "Compacte a pasta do projeto em .zip — SEM a pasta node_modules.",
        "Nomeie o arquivo como: SeuNome-Aula04-Login.zip",
        "Envie até: _______________________",
        "Canal de entrega: _______________________",
    ])
    caixa(doc, "Por que remover a node_modules",
          "Ela costuma ter centenas de megabytes e é recriada com um "
          "npm install. Mandar o projeto sem ela é o padrão do mercado.")

    # ---------------------------------------------------------- 8
    h1(doc, "8. Critérios de avaliação")
    criterios = [
        ("O projeto roda sem erros com npm run dev", "20"),
        ("O login valida e-mail/senha e mostra a mensagem de erro", "20"),
        ("O botão mostrar/ocultar senha funciona corretamente", "25"),
        ("O CSS é autoral, diferente do exemplo da aula", "20"),
        ("A segunda página foi iniciada", "15"),
        ("Total", "100"),
    ]
    tabela = doc.add_table(rows=len(criterios) + 1, cols=2)
    tabela.style = "Table Grid"
    cab = tabela.rows[0].cells
    for i, texto in enumerate(("Critério", "Pontos")):
        set_cell_bg(cab[i], "4F46E5")
        run = cab[i].paragraphs[0].add_run(texto)
        run.font.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = SANS
    for linha, (criterio, pontos) in enumerate(criterios, start=1):
        celulas = tabela.rows[linha].cells
        ultimo = criterio == "Total"
        if ultimo:
            set_cell_bg(celulas[0], "EEF0FE")
            set_cell_bg(celulas[1], "EEF0FE")
        for celula, texto in zip(celulas, (criterio, pontos)):
            run = celula.paragraphs[0].add_run(texto)
            run.font.size = Pt(10.5)
            run.font.bold = ultimo
            run.font.name = SANS
        celulas[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragrafo(doc, "", size=6, space_after=8)

    # ---------------------------------------------------------- 9
    h1(doc, "9. Desafio extra (opcional)")
    paragrafo(doc,
              "Terminou antes? Estes são bons próximos passos — não valem "
              "nota, valem aprendizado:")
    lista(doc, [
        [("Validar o formato do e-mail antes de aceitar (procure por ",
          False), ("includes(\"@\")", False, True), (").", False)],
        "Desabilitar o botão Entrar enquanto os campos estiverem vazios.",
        "Trocar o texto Mostrar/Ocultar por um ícone de olho.",
        "Mostrar a mensagem de erro com uma transição suave em vez de "
        "aparecer de uma vez.",
        "Guardar o e-mail digitado para que ele continue lá ao recarregar a "
        "página.",
    ])


def build(destino):
    doc = Document()
    estilo_base(doc)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    rodape_numerado(section)

    atualizar_campos_ao_abrir(doc)
    capa(doc)
    sumario(doc)
    conteudo(doc)

    doc.save(destino)
    print(f"OK: {destino}")


if __name__ == "__main__":
    saida = Path(__file__).with_name("Trabalho-Login-Aula04.docx")
    build(str(saida))
